"""In-process tests for intent scoring, lead rollup, admin reports/exports,
staff roles and the content CMS (mongomock-motor, no external services)."""
import io
import os

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "solix_test")
os.environ.setdefault("JWT_SECRET", "test-secret")
os.environ.setdefault("ADMIN_EMAIL", "admin@example.com")
os.environ.setdefault("ADMIN_PASSWORD", "AdminPass!1")

import json  # noqa: E402
from datetime import datetime, timedelta, timezone  # noqa: E402

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from mongomock_motor import AsyncMongoMockClient  # noqa: E402

import database  # noqa: E402

import accounts, admin, auth, cache, chat, content, emailer, intent, leads_admin, press, server  # noqa: E402,E401

MODULES = (accounts, admin, auth, chat, content, emailer, intent, leads_admin, press, server)
PDF = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF"
PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 32


@pytest.fixture()
def client():
    mock = AsyncMongoMockClient()["solix_test"]
    previous = {m: getattr(m, "db", None) for m in MODULES}
    previous_database = database.db
    database.db = mock
    for m in MODULES:
        m.db = mock
    cache.clear()
    with TestClient(server.app) as c:
        yield c
    # Put back whatever database other test modules patched in.
    database.db = previous_database
    for m, db in previous.items():
        m.db = db


def staff(client, email="admin@example.com", password="AdminPass!1"):
    r = client.post("/api/auth/login", json={"email": email, "password": password})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def track(client, vid, events, **ctx):
    return client.post("/api/track", content=json.dumps({"visitor_id": vid, "session_id": "s1", "events": events, "context": ctx}), headers={"Content-Type": "text/plain"})


# --- Intent scoring -------------------------------------------------------------

def test_tracking_scores_products_and_ignores_client_weights(client):
    vid = "visitor-0001"
    r = track(client, vid, [
        {"type": "page_view", "path": "/products/sap-archiving", "points": 999},
        {"type": "engaged", "path": "/products/sap-archiving"},
        {"type": "page_view", "path": "/products/sap-archiving"},  # repeat in same session: no extra points
        {"type": "resource_download", "path": "/resources/x", "topics": ["enterprise-archiving", "not-a-product"]},
        {"type": "bogus", "path": "/"},
    ], utm_source="google", utm_medium="cpc", referrer="https://www.google.com/")
    assert r.status_code == 204
    import asyncio
    v = asyncio.get_event_loop().run_until_complete(intent.db.visitors.find_one({"id": vid}, {"_id": 0}))
    assert v["product_scores"] == {"sap-archiving": 8.0, "enterprise-archiving": 15.0}
    assert v["first_touch"]["utm_medium"] == "cpc"
    assert v["pages_count"] == 2


def test_form_links_visitor_and_tags_lead_to_product_line(client):
    vid = "visitor-ai-01"
    track(client, vid, [{"type": "page_view", "path": "/products/data-ask"}, {"type": "engaged", "path": "/products/data-ask"}, {"type": "chat_topic", "path": "/", "topics": ["agentic"]}])
    r = client.post("/api/submissions", json={"type": "download", "email": "Ana@Bank.com", "name": "Ana", "company": "Bank", "job_title": "VP Data", "visitor_id": vid, "topics": ["enterprise-ai"]})
    assert r.status_code == 201
    h = staff(client)
    leads = client.get("/api/admin/leads", headers=h).json()
    assert leads["total"] == 1
    lead = leads["items"][0]
    assert lead["email"] == "ana@bank.com"
    assert lead["primary_line"] == "ai" and lead["primary_line_label"] == "Enterprise AI"
    assert lead["fit_score"] == 25  # business email + seniority
    # 3 + 5 + 6 behaviour + 15 download = 29 on AI, + 25 fit = 54 >= 45 threshold
    assert lead["score"] == 54.0 and lead["stage"] == "mql"
    assert "threshold" in lead["stage_reason"]
    # Later browsing by the same visitor keeps rolling up into the lead.
    track(client, vid, [{"type": "page_view", "path": "/products/ai-governance"}])
    detail = client.get(f"/api/admin/leads/{lead['id']}", headers=h).json()
    assert detail["lead"]["product_scores"]["ai-governance"] == 3.0
    kinds = {t["kind"] for t in detail["timeline"]}
    assert kinds == {"form", "event", "stage"}


def test_hand_raiser_is_mql_and_free_email_scores_low_fit(client):
    client.post("/api/submissions", json={"type": "demo", "email": "joe@gmail.com", "name": "Joe", "interest": "enterprise-archiving"})
    lead = client.get("/api/admin/leads", headers=staff(client)).json()["items"][0]
    assert lead["stage"] == "mql" and "Hand-raiser" in lead["stage_reason"]
    assert lead["fit_score"] == 0
    assert lead["primary_line"] == "archiving"


def test_careers_forms_do_not_create_leads(client):
    client.post("/api/submissions", json={"type": "career", "email": "dev@x.com", "name": "Dev"})
    assert client.get("/api/admin/leads", headers=staff(client)).json()["total"] == 0


def test_trial_signup_and_chat_booking_create_leads(client):
    r = client.post("/api/accounts/signup", json={"first_name": "Li", "last_name": "Wu", "company": "Acme", "email": "li@acme.io", "phone": "+1 555", "password": "Engine#42", "company_size": "1,000-4,999", "visitor_id": "no-such-visitor"})
    assert r.status_code == 201
    import asyncio
    asyncio.get_event_loop().run_until_complete(chat.create_demo_request("sess-1", {"name": "Kim Lee", "email": "kim@corp.com", "company": "Corp"}))
    items = {l["email"]: l for l in client.get("/api/admin/leads", headers=staff(client)).json()["items"]}
    assert items["li@acme.io"]["primary_line"] == "ecs" and items["li@acme.io"]["stage"] == "mql"
    assert items["li@acme.io"]["fit_score"] == 25  # business email + 1,000+ + phone
    assert items["kim@corp.com"]["channel"] == "chat"


def test_filters_patch_bulk_views_and_export(client):
    for i, (t, interest) in enumerate([("demo", "enterprise-archiving"), ("newsletter", "enterprise-ai"), ("contact", "consumer-data-privacy")]):
        client.post("/api/submissions", json={"type": t, "email": f"p{i}@co{i}.com", "name": f"P{i}", "country": "Germany" if i else "India", "interest": interest})
    h = staff(client)
    assert client.get("/api/admin/leads?line=archiving", headers=h).json()["total"] == 1
    assert client.get("/api/admin/leads?stage=mql", headers=h).json()["total"] == 2
    assert client.get("/api/admin/leads?country=Germany", headers=h).json()["total"] == 2
    lead = client.get("/api/admin/leads?line=ai", headers=h).json()["items"][0]
    r = client.patch(f"/api/admin/leads/{lead['id']}", json={"stage": "sql", "owner": "Rep@Solix.com", "notes": "Budget confirmed"}, headers=h)
    assert r.status_code == 200 and r.json()["stage"] == "sql" and r.json()["owner"] == "rep@solix.com"
    ids = [l["id"] for l in client.get("/api/admin/leads", headers=h).json()["items"]]
    assert client.post("/api/admin/leads/bulk", json={"ids": ids, "owner": "team@solix.com"}, headers=h).json()["updated"] == 3
    assert client.get("/api/admin/leads?owner=team@solix.com", headers=h).json()["total"] == 3
    view = client.post("/api/admin/views", json={"name": "EU MQLs", "filters": {"country": "Germany", "stage": "mql", "junk": "x"}}, headers=h).json()
    assert view["filters"] == {"country": "Germany", "stage": "mql"}
    csv_text = client.get("/api/admin/leads-export?format=csv", headers=h).text
    assert csv_text.splitlines()[0].startswith("Name,Email,Company") and len(csv_text.splitlines()) == 4
    x = client.get("/api/admin/leads-export?format=xlsx&stage=sql", headers=h)
    assert x.status_code == 200 and x.content[:2] == b"PK"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Leads", "Summary by product line"] and wb["Leads"].max_row == 2


def test_export_neutralises_formula_injection(client):
    client.post("/api/submissions", json={"type": "contact", "email": "evil@x.com", "name": "=HYPERLINK(\"http://x\")"})
    csv_text = client.get("/api/admin/leads-export", headers=staff(client)).text
    assert "'=HYPERLINK" in csv_text


def test_overview_report(client):
    track(client, "visitor-rep-1", [{"type": "resource_view", "path": "/resources/cloud-archive", "topics": ["enterprise-archiving"], "meta": {"title": "Cloud archive"}}, {"type": "resource_download", "path": "/resources/cloud-archive", "topics": ["enterprise-archiving"]}])
    client.post("/api/submissions", json={"type": "demo", "email": "a@corp.com", "visitor_id": "visitor-rep-1", "interest": "sap-archiving"})
    client.post("/api/submissions", json={"type": "newsletter", "email": "b@corp.com"})
    r = client.get("/api/admin/reports/overview?days=7", headers=staff(client))
    assert r.status_code == 200
    d = r.json()
    assert d["kpis"]["leads"] == 2 and d["kpis"]["mqls"] == 1 and d["kpis"]["new_visitors"] == 1
    assert d["kpis"]["lead_to_mql"] == 50.0
    assert len(d["trend"]) == 7 and d["trend"][-1]["leads"] == 2
    arch = next(r for r in d["by_line"] if r["key"] == "archiving")
    assert arch["leads"] == 1 and arch["mqls"] == 1
    assert d["top_content"][0] == {"path": "/resources/cloud-archive", "title": "Cloud archive", "views": 1, "downloads": 1}
    assert d["funnel"][0] == {"stage": "visitors", "count": 1}


def test_scoring_settings_rescore(client):
    client.post("/api/submissions", json={"type": "download", "email": "c@corp.com", "interest": "ediscovery"})
    h = staff(client)
    assert client.get("/api/admin/leads", headers=h).json()["items"][0]["stage"] == "lead"  # 15 + 10 fit < 45
    r = client.put("/api/admin/scoring", json={"mql_threshold": 20, "half_life_days": 30}, headers=h)
    assert r.status_code == 200 and r.json()["rescored"] == 1
    assert client.get("/api/admin/leads", headers=h).json()["items"][0]["stage"] == "mql"


# --- Roles ----------------------------------------------------------------------

def test_roles(client):
    h = staff(client)
    created = client.post("/api/admin/users", json={"email": "ceo@solix.com", "name": "CEO", "role": "viewer"}, headers=h).json()
    viewer = staff(client, "ceo@solix.com", created["temporary_password"])
    client.post("/api/submissions", json={"type": "demo", "email": "z@corp.com"})
    lead = client.get("/api/admin/leads", headers=viewer).json()["items"][0]  # can read
    assert client.get("/api/admin/reports/overview", headers=viewer).status_code == 200
    assert client.patch(f"/api/admin/leads/{lead['id']}", json={"stage": "won"}, headers=viewer).status_code == 403
    assert client.post("/api/admin/content", json={"title": "Hello", "type": "blog"}, headers=viewer).status_code == 403
    assert client.get("/api/admin/users", headers=viewer).status_code == 403
    client.patch(f"/api/admin/users/{created['user']['id']}", json={"disabled": True}, headers=h)
    assert client.get("/api/admin/leads", headers=viewer).status_code == 401
    assert client.post("/api/auth/login", json={"email": "ceo@solix.com", "password": created["temporary_password"]}).status_code == 401


def test_password_change(client):
    h = staff(client)
    assert client.post("/api/auth/password", json={"current_password": "wrong", "new_password": "NewPassword123"}, headers=h).status_code == 400
    assert client.post("/api/auth/password", json={"current_password": "AdminPass!1", "new_password": "short"}, headers=h).status_code == 422


# --- CMS --------------------------------------------------------------------------

def test_publish_flow_caching_and_versions(client):
    h = staff(client)
    r = client.post("/api/admin/content", json={"title": "Archive-first S/4HANA", "type": "blog", "summary": "Why.", "body": "## Intro\n\nText " * 50, "products": ["sap-archiving", "nope"]}, headers=h)
    assert r.status_code == 201
    item = r.json()
    assert item["slug"] == "archive-first-s-4hana" and item["products"] == ["sap-archiving"] and item["status"] == "draft"
    assert client.get("/api/content").json()["items"] == []  # drafts are private
    client.post(f"/api/admin/content/{item['id']}/publish", json={}, headers=h)
    r = client.get("/api/content")
    assert r.status_code == 200 and "max-age=60" in r.headers["cache-control"]
    assert [i["slug"] for i in r.json()["items"]] == ["archive-first-s-4hana"] and "body" not in r.json()["items"][0]
    assert client.get("/api/content", headers={"If-None-Match": r.headers["etag"]}).status_code == 304
    full = client.get("/api/content?full=true").json()["items"][0]
    assert full["body"].startswith("## Intro")
    assert client.get("/api/content/archive-first-s-4hana").json()["read_minutes"] >= 1
    # Edit -> version saved; restore brings the old title back.
    client.put(f"/api/admin/content/{item['id']}", json={"title": "New title", "type": "blog", "slug": "archive-first-s-4hana"}, headers=h)
    assert client.get("/api/content/archive-first-s-4hana").json()["title"] == "New title"
    versions = client.get(f"/api/admin/content/{item['id']}/versions", headers=h).json()
    assert versions[0]["title"] == "Archive-first S/4HANA"
    client.post(f"/api/admin/content/{item['id']}/versions/{versions[0]['version']}/restore", headers=h)
    assert client.get("/api/content/archive-first-s-4hana").json()["title"] == "Archive-first S/4HANA"
    # Duplicate slug rejected.
    assert client.post("/api/admin/content", json={"title": "Archive-first S/4HANA", "type": "blog"}, headers=h).status_code == 409
    client.post(f"/api/admin/content/{item['id']}/unpublish", headers=h)
    assert client.get("/api/content").json()["items"] == []


def test_scheduled_content_is_hidden_until_due(client):
    h = staff(client)
    item = client.post("/api/admin/content", json={"title": "Future post", "type": "blog"}, headers=h).json()
    future = (datetime.now(timezone.utc) + timedelta(days=2)).isoformat()
    assert client.post(f"/api/admin/content/{item['id']}/publish", json={"publish_at": future}, headers=h).json()["status"] == "scheduled"
    assert client.get("/api/content").json()["items"] == []


def test_files_gating_and_type_checks(client):
    h = staff(client)
    assert client.post("/api/admin/files", files={"file": ("x.svg", b"<svg onload=alert(1)>", "image/svg+xml")}, headers=h).status_code == 415
    img = client.post("/api/admin/files", files={"file": ("Cover Photo.PNG", PNG, "image/png")}, headers=h).json()
    assert img["kind"] == "image" and img["name"] == "cover-photo.png"
    r = client.get(img["public_url"])
    assert r.status_code == 200 and "immutable" in r.headers["cache-control"] and r.headers["x-content-type-options"] == "nosniff"
    pdf = client.post("/api/admin/files", files={"file": ("Datasheet.pdf", PDF, "application/pdf")}, headers=h).json()
    assert client.get(pdf["public_url"]).status_code == 403  # not attached to anything public yet
    item = client.post("/api/admin/content", json={"title": "ECS datasheet", "type": "datasheet", "file_id": pdf["id"], "gated": True}, headers=h).json()
    client.post(f"/api/admin/content/{item['id']}/publish", json={}, headers=h)
    pub = client.get("/api/content/ecs-datasheet").json()
    assert pub["file"]["gated"] is True and "url" not in pub["file"]
    assert client.get(pdf["public_url"]).status_code == 403
    assert client.post("/api/content/ecs-datasheet/unlock", json={"submission_id": "made-up-id"}).status_code == 403
    sub = client.post("/api/submissions", json={"type": "download", "email": "r@corp.com", "resource": "ECS datasheet"}).json()
    url = client.post("/api/content/ecs-datasheet/unlock", json={"submission_id": sub["id"]}).json()["url"]
    got = client.get(url)
    assert got.status_code == 200 and got.content == PDF and got.headers["content-type"] == "application/pdf"
    assert client.get(url.replace("?t=", "?t=9") ).status_code == 403
    # Un-gated attachment becomes public.
    client.put(f"/api/admin/content/{item['id']}", json={"title": "ECS datasheet", "type": "datasheet", "file_id": pdf["id"], "gated": False}, headers=h)
    cache.bump("content")
    assert client.get("/api/content/ecs-datasheet").json()["file"]["url"] == pdf["public_url"]
    assert client.get(pdf["public_url"]).status_code == 200


def test_track_guards(client):
    assert track(client, "short", [{"type": "page_view", "path": "/"}]).status_code == 204
    assert client.post("/api/track", content=b"not json").status_code == 400
    assert client.post("/api/track", content=b"x" * 40000).status_code == 413


def test_form_before_first_beacon_still_links_browsing_and_campaign(client):
    vid = "visitor-race-1"
    client.post("/api/submissions", json={"type": "download", "email": "sam@corp.com", "visitor_id": vid})
    track(client, vid, [{"type": "page_view", "path": "/products/enterprise-ai"}], utm_source="linkedin", utm_medium="paid_social", landing="/products/enterprise-ai")
    lead = client.get("/api/admin/leads", headers=staff(client)).json()["items"][0]
    assert lead["product_scores"] == {"enterprise-ai": 3.0}
    assert lead["first_touch"]["utm_source"] == "linkedin" and lead["channel"] == "paid"
