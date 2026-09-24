"""In-process tests for event registration, tickets, payments and admin (mongomock)."""
import os

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "solix_test")
os.environ.setdefault("JWT_SECRET", "test-secret")
os.environ.setdefault("ADMIN_EMAIL", "admin@example.com")
os.environ.setdefault("ADMIN_PASSWORD", "AdminPass!1")

import io  # noqa: E402

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from mongomock_motor import AsyncMongoMockClient  # noqa: E402

import database  # noqa: E402
import accounts, admin, auth, cache, chat, content, emailer, events, intent, leads_admin, server  # noqa: E402,E401

MODULES = (accounts, admin, auth, chat, content, emailer, events, intent, leads_admin, server)
SLUG = "empower-2026"


@pytest.fixture()
def client():
    mock = AsyncMongoMockClient()["solix_test"]
    previous = {m: getattr(m, "db", None) for m in MODULES}
    previous_database = database.db
    database.db = mock
    for m in MODULES:
        m.db = mock
    cache.clear()
    events._hits.clear()
    with TestClient(server.app) as c:
        yield c
    database.db = previous_database
    for m, db in previous.items():
        m.db = db


def staff(client):
    r = client.post("/api/auth/login", json={"email": "admin@example.com", "password": "AdminPass!1"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def reg(client, **over):
    body = {"ticket_id": "full-pass", "first_name": "Ana", "last_name": "Diaz", "email": "ana@bank.com", "company": "Big Bank", "job_title": "VP Data",
            "interests": ["enterprise-ai", "data-governance", "bogus"], "days": ["day1", "day2", "day9"], "dinners": ["day1"], "accept_terms": True,
            "utm": {"utm_source": "solix_website", "utm_medium": "promo_banner"}}
    body.update(over)
    return client.post(f"/api/events/{SLUG}/registrations", json=body)


def test_public_event_is_seeded(client):
    d = client.get(f"/api/events/{SLUG}").json()
    assert d["name"] == "SOLIXEmpower 2026" and d["registration_open"] is True
    assert [t["id"] for t in d["tickets"]] == ["full-pass"] and d["tickets"][0]["price"] == 0
    assert "promo_codes" not in d and len(d["interests"]) == 6


def test_free_registration_confirms_and_scores_lead(client):
    r = reg(client)
    assert r.status_code == 201
    j = r.json()
    assert j["status"] == "confirmed" and j["code"].startswith("EMP-") and j["payment"] is None
    assert j["interests"] == ["enterprise-ai", "data-governance"] and j["days"] == ["day1", "day2"]
    # Same email again returns the existing registration instead of a duplicate.
    again = reg(client).json()
    assert again["already_registered"] and again["code"] == j["code"]
    h = staff(client)
    lead = client.get("/api/admin/leads", headers=h).json()["items"][0]
    assert lead["email"] == "ana@bank.com" and "empower-2026" in lead["tags"] and lead["primary_line"] == "ai"
    assert client.get("/api/admin/submissions?type=event", headers=h).json()["total"] == 1
    # Confirmation lookup needs the matching email.
    assert client.post(f"/api/events/{SLUG}/registrations/{j['code']}/lookup", json={"email": "other@x.com"}).status_code == 404
    assert client.post(f"/api/events/{SLUG}/registrations/{j['code']}/lookup", json={"email": "ANA@bank.com"}).json()["status"] == "confirmed"


def test_terms_and_closed_registration(client):
    assert reg(client, accept_terms=False).status_code == 422
    h = staff(client)
    ev = client.get(f"/api/admin/events/{SLUG}", headers=h).json()
    ev["registration_open"] = False
    assert client.put(f"/api/admin/events/{SLUG}", json=ev, headers=h).status_code == 200
    assert reg(client).status_code == 409


def paid_setup(client, **ticket):
    h = staff(client)
    ev = client.get(f"/api/admin/events/{SLUG}", headers=h).json()
    ev["tickets"].append({"id": "workshop", "name": "Hands-on workshop", "price": 49900, "currency": "USD", "active": True, **ticket})
    ev["promo_codes"] = [{"code": "slxemp26eb", "percent_off": 20, "active": True, "max_uses": 1}]
    r = client.put(f"/api/admin/events/{SLUG}", json=ev, headers=h)
    assert r.status_code == 200, r.text
    return h


def test_stripe_link_payment_with_promo(client):
    paid_setup(client, provider="stripe_link", payment_link="https://buy.stripe.com/test_123")
    q = client.post(f"/api/events/{SLUG}/quote", json={"ticket_id": "workshop", "promo_code": "SLXEMP26EB"}).json()
    assert q == {"price": 49900, "discount": 9980, "total": 39920, "currency": "USD", "promo_valid": True, "promo_message": "20% off applied"}
    j = reg(client, ticket_id="workshop", promo_code="slxemp26eb").json()
    assert j["status"] == "pending_payment" and j["amount"] == 39920
    assert j["payment"]["url"].startswith("https://buy.stripe.com/test_123?client_reference_id=EMP-")
    # Promo max_uses=1 is now spent.
    assert client.post(f"/api/events/{SLUG}/quote", json={"ticket_id": "workshop", "promo_code": "SLXEMP26EB"}).json()["promo_valid"] is False
    r = client.post(f"/api/events/{SLUG}/registrations/{j['code']}/payment", json={"email": "ana@bank.com", "provider": "stripe_link", "reference": "cs_test"})
    assert r.json()["status"] == "payment_reported"


def test_eventbrite_and_invoice_providers(client):
    paid_setup(client, provider="eventbrite")
    j = reg(client, ticket_id="workshop").json()
    assert j["payment"] == {"provider": "eventbrite", "eventbrite_event_id": "1994300379119", "promo_code": None}
    h = staff(client)
    ev = client.get(f"/api/admin/events/{SLUG}", headers=h).json()
    ev["tickets"][1]["provider"] = "invoice"
    client.put(f"/api/admin/events/{SLUG}", json=ev, headers=h)
    assert reg(client, email="bo@corp.com", ticket_id="workshop", po_number="PO-1").json()["status"] == "invoice_requested"


def test_settings_validation(client):
    h = staff(client)
    ev = client.get(f"/api/admin/events/{SLUG}", headers=h).json()
    ev["tickets"].append({"id": "vip", "name": "VIP", "price": 1000, "provider": "stripe_link", "active": True})
    assert "Stripe Payment Link" in client.put(f"/api/admin/events/{SLUG}", json=ev, headers=h).json()["detail"]
    ev["tickets"][-1] = {"id": "vip", "name": "VIP", "price": 0, "provider": "invoice", "active": True}
    assert "no price" in client.put(f"/api/admin/events/{SLUG}", json=ev, headers=h).json()["detail"]


def test_capacity_waitlist_admin_stats_and_export(client):
    h = staff(client)
    ev = client.get(f"/api/admin/events/{SLUG}", headers=h).json()
    ev["capacity"] = 1
    client.put(f"/api/admin/events/{SLUG}", json=ev, headers=h)
    assert reg(client).json()["status"] == "confirmed"
    assert reg(client, email="late@corp.com").json()["status"] == "waitlisted"
    assert client.get(f"/api/events/{SLUG}").json()["waitlist"] is True
    d = client.get(f"/api/admin/events/{SLUG}/registrations", headers=h).json()
    assert d["total"] == 2 and d["stats"]["registered"] == 1 and d["stats"]["by_status"]["waitlisted"] == 1
    assert d["stats"]["by_interest"] == {"data-governance": 1, "enterprise-ai": 1} and d["stats"]["sources"][0] == {"source": "solix_website", "count": 1}
    rid = next(i["id"] for i in d["items"] if i["status"] == "confirmed")
    assert client.patch(f"/api/admin/events/{SLUG}/registrations/{rid}", json={"checked_in": True}, headers=h).json()["checked_in"] is True
    assert client.get(f"/api/admin/events/{SLUG}/registrations?checked_in=true", headers=h).json()["total"] == 1
    csv_text = client.get(f"/api/admin/events/{SLUG}/registrations-export", headers=h).text
    assert csv_text.splitlines()[0].startswith("Code,Status,Pass") and len(csv_text.splitlines()) == 3
    x = client.get(f"/api/admin/events/{SLUG}/registrations-export?format=xlsx", headers=h)
    from openpyxl import load_workbook
    assert load_workbook(io.BytesIO(x.content))["Registrations"].max_row == 3
    assert client.get("/api/admin/events", headers=h).json()[0]["registered"] == 1
