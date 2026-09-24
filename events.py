"""Events: SOLIXEmpower (and future events) registration, tickets and payments.

The Empower website (a separate site) reads public event settings and posts
registrations here. Each ticket chooses how it is paid for:

  free        confirmed immediately (complimentary passes)
  eventbrite  checkout in the Eventbrite widget (Eventbrite takes payment);
              the site reports the order id back and staff can verify it
  stripe_link redirect to a Stripe Payment Link (no API keys needed); the
              registration code travels as client_reference_id
  invoice     the registration is held and the events team invoices the
              company (PO / bank transfer), then marks it paid

Every registration is also a sales signal: it lands in the form inbox, fires
the usual lead alert, and is scored into the attendee's lead (see intent.py)
using the tracks they chose.
"""
from __future__ import annotations

import asyncio
import csv
import io
import logging
import re
import secrets
import time
import uuid
from collections import defaultdict, deque
from datetime import datetime, timezone
from html import escape
from typing import Dict, List, Literal, Optional

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, EmailStr, Field

import cache
from auth import get_current_admin, require_roles
from database import db, now_iso
from emailer import email_provider, notify_lead, send_email
from intent import record_submission

logger = logging.getLogger("solix.events")
public = APIRouter(prefix="/api/events", tags=["events"])
admin = APIRouter(prefix="/api/admin/events", tags=["events-admin"], dependencies=[Depends(get_current_admin)])
can_configure = Depends(require_roles("admin", "editor"))
can_work = Depends(require_roles("admin", "sales", "editor"))

Provider = Literal["free", "eventbrite", "stripe_link", "invoice"]
STATUSES = ("confirmed", "pending_payment", "payment_reported", "paid", "invoice_requested", "waitlisted", "cancelled")
ACTIVE = ("confirmed", "pending_payment", "payment_reported", "paid", "invoice_requested")

# Registration interests -> product slugs used by lead scoring.
INTERESTS: Dict[str, dict] = {
    "enterprise-ai": {"label": "Enterprise AI & agents", "topics": ["enterprise-ai", "agentic"]},
    "data-governance": {"label": "Enterprise Data Governance", "topics": ["enterprise-data-governance", "ai-governance"]},
    "cloud-data-management": {"label": "Cloud Data Management", "topics": ["common-data-platform", "enterprise-archiving"]},
    "knowledge-graph": {"label": "Application Knowledge Graph (AKG)", "topics": ["application-knowledge-graph", "data-sense"]},
    "life-sciences": {"label": "AI in Pharma & Life Sciences", "topics": ["eai-pharma", "ai-healthcare"]},
    "content-services": {"label": "Content Intelligence (ECS)", "topics": ["enterprise-content-services"]},
}

DEFAULT_EVENTS = {
    "empower-2026": {
        "slug": "empower-2026",
        "name": "SOLIXEmpower 2026",
        "theme": "The Agentic Enterprise: Reimagining Enterprise Applications with Enterprise AI",
        "starts_at": "2026-10-28T07:30:00-07:00",
        "ends_at": "2026-10-30T13:30:00-07:00",
        "timezone": "America/Los_Angeles",
        "venue": "The Qualcomm Institute, Atkinson Hall, UC San Diego",
        "address": "3195 Voigt Drive, La Jolla, CA 92093",
        "registration_open": True,
        "capacity": None,
        "contact_email": "info@solixempower.com",
        "site_url": "https://akshatsingh-solix.github.io/Website/empower/",
        "eventbrite_event_id": "1994300379119",
        "days": [
            {"id": "day1", "label": "Oct 28 · Day 1", "dinner": "Solix User Group Cocktails & Dinner"},
            {"id": "day2", "label": "Oct 29 · Day 2", "dinner": "Dinner & San Diego Supercomputer Center tour"},
            {"id": "day3", "label": "Oct 30 · Day 3 (half day)", "dinner": None},
        ],
        "tickets": [
            {
                "id": "full-pass", "name": "Full Event Pass", "active": True, "price": 29900, "currency": "USD", "provider": "eventbrite",
                "description": "All three days, Oct 28-30: keynotes, panels, hands-on workshops, the hackathon finals, networking meals and evening receptions.",
                "capacity": None, "eventbrite_event_id": None, "payment_link": None, "sales_end_at": "2026-10-28T23:59:00-07:00",
            },
        ],
        "refund_policy": "Refunds up to 7 days before the event. Eventbrite's fee is non-refundable.",
        "promo_codes": [],
        "seed_version": 2,
    }
}
# Seed fields that a newer seed_version may update on events staff haven't edited yet.
SEED_FIELDS = ("tickets", "refund_policy")

CODE_ALPHABET = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"
_hits: dict = defaultdict(deque)


def _rate_limited(key: str, limit: int = 12, window: int = 60) -> bool:
    now = time.monotonic()
    q = _hits[key]
    while q and now - q[0] > window:
        q.popleft()
    if len(q) >= limit:
        return True
    q.append(now)
    return False


def _ip(request: Request) -> str:
    fwd = request.headers.get("x-forwarded-for", "")
    return fwd.split(",")[0].strip() if fwd else (request.client.host if request.client else "unknown")


def _new_code() -> str:
    return "EMP-" + "".join(secrets.choice(CODE_ALPHABET) for _ in range(6))


async def get_event(slug: str) -> dict:
    doc = await db.event_configs.find_one({"slug": slug}, {"_id": 0})
    seed = DEFAULT_EVENTS.get(slug)
    if doc and seed and (doc.get("seed_version") or 1) < seed["seed_version"]:
        # Bring an untouched seeded event up to date (e.g. the pass became paid);
        # anything staff have saved from the admin is left alone.
        patch = {"seed_version": seed["seed_version"], "updated_at": now_iso()}
        if not doc.get("updated_by"):
            patch.update({k: seed[k] for k in SEED_FIELDS})
        await db.event_configs.update_one({"slug": slug}, {"$set": patch})
        cache.bump("events")
        doc = await db.event_configs.find_one({"slug": slug}, {"_id": 0})
    if doc:
        return doc
    if slug in DEFAULT_EVENTS:
        doc = {**DEFAULT_EVENTS[slug], "created_at": now_iso(), "updated_at": now_iso()}
        await db.event_configs.update_one({"slug": slug}, {"$setOnInsert": doc}, upsert=True)
        return await db.event_configs.find_one({"slug": slug}, {"_id": 0})
    raise HTTPException(status_code=404, detail="Event not found")


def _price(ticket: dict, promo: Optional[dict]) -> dict:
    base = int(ticket.get("price") or 0)
    discount = round(base * (promo["percent_off"] / 100)) if promo else 0
    return {"price": base, "discount": discount, "total": max(0, base - discount), "currency": ticket.get("currency") or "USD"}


def _sales_ended(ticket: dict) -> bool:
    end = ticket.get("sales_end_at")
    if not end:
        return False
    try:
        return datetime.now(timezone.utc) > datetime.fromisoformat(end)
    except ValueError:
        return False


def _find_promo(event: dict, code: Optional[str], ticket_id: str) -> Optional[dict]:
    if not code:
        return None
    code = code.strip().upper()
    for p in event.get("promo_codes") or []:
        if p.get("code", "").upper() == code and p.get("active"):
            if p.get("tickets") and ticket_id not in p["tickets"]:
                return None
            if p.get("max_uses") and (p.get("uses") or 0) >= p["max_uses"]:
                return None
            return p
    return None


async def _counts(slug: str) -> Dict[str, int]:
    counts: Dict[str, int] = {}
    async for r in db.event_registrations.find({"event": slug, "status": {"$in": list(ACTIVE)}}, {"_id": 0, "ticket_id": 1}):
        counts[r["ticket_id"]] = counts.get(r["ticket_id"], 0) + 1
    counts["_total"] = sum(v for k, v in counts.items())
    return counts


def _public_ticket(t: dict, sold: int, event: dict) -> dict:
    cap = t.get("capacity")
    out = {k: t.get(k) for k in ("id", "name", "description", "price", "currency", "provider", "sales_end_at")}
    out["seats_left"] = max(0, cap - sold) if cap else None
    out["sold_out"] = bool(cap and sold >= cap)
    out["sales_ended"] = _sales_ended(t)
    if t.get("provider") == "eventbrite":
        out["eventbrite_event_id"] = t.get("eventbrite_event_id") or event.get("eventbrite_event_id")
    return out


# --- Public -----------------------------------------------------------------

@public.get("/{slug}")
async def public_event(slug: str, request: Request):
    async def produce():
        event = await get_event(slug)
        counts = await _counts(slug)
        cap = event.get("capacity")
        tickets = [_public_ticket(t, counts.get(t["id"], 0), event) for t in event.get("tickets") or [] if t.get("active")]
        return {
            **{k: event.get(k) for k in ("slug", "name", "theme", "starts_at", "ends_at", "timezone", "venue", "address", "contact_email", "days", "refund_policy")},
            "registration_open": bool(event.get("registration_open")),
            "seats_left": max(0, cap - counts["_total"]) if cap else None,
            "waitlist": bool(cap and counts["_total"] >= cap),
            "tickets": tickets,
            "interests": [{"key": k, "label": v["label"]} for k, v in INTERESTS.items()],
            "has_promo_codes": any(p.get("active") for p in event.get("promo_codes") or []),
        }
    data = await cache.memo("events", slug, 20, produce)
    return cache.cached_json(request, data, max_age=30, swr=600)


class QuoteIn(BaseModel):
    ticket_id: str = Field(max_length=60)
    promo_code: Optional[str] = Field(default=None, max_length=40)


@public.post("/{slug}/quote")
async def quote(slug: str, body: QuoteIn, request: Request):
    if _rate_limited(f"quote:{_ip(request)}", 30):
        raise HTTPException(status_code=429, detail="Too many attempts. Try again in a minute.")
    event = await get_event(slug)
    ticket = next((t for t in event.get("tickets") or [] if t["id"] == body.ticket_id and t.get("active")), None)
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")
    promo = _find_promo(event, body.promo_code, ticket["id"])
    if not promo and body.promo_code and ticket.get("provider") == "eventbrite":
        # Eventbrite checks its own discount codes at checkout.
        return {**_price(ticket, None), "promo_valid": None, "promo_message": "Your code will be applied at Eventbrite checkout."}
    return {**_price(ticket, promo), "promo_valid": bool(promo), "promo_message": (f"{promo['percent_off']}% off applied" if promo else ("That code isn't valid for this pass." if body.promo_code else None))}


class RegistrationIn(BaseModel):
    ticket_id: str = Field(max_length=60)
    first_name: str = Field(min_length=1, max_length=80)
    last_name: str = Field(min_length=1, max_length=80)
    email: EmailStr
    company: str = Field(min_length=1, max_length=160)
    job_title: str = Field(min_length=1, max_length=120)
    phone: Optional[str] = Field(default=None, max_length=40)
    country: Optional[str] = Field(default=None, max_length=80)
    interests: List[str] = Field(default_factory=list, max_length=10)
    days: List[str] = Field(default_factory=list, max_length=5)
    dinners: List[str] = Field(default_factory=list, max_length=5)
    hackathon: bool = False
    dietary: Optional[str] = Field(default=None, max_length=200)
    accessibility: Optional[str] = Field(default=None, max_length=300)
    how_heard: Optional[str] = Field(default=None, max_length=80)
    promo_code: Optional[str] = Field(default=None, max_length=40)
    billing_contact: Optional[str] = Field(default=None, max_length=160)
    po_number: Optional[str] = Field(default=None, max_length=60)
    marketing_opt_in: bool = False
    accept_terms: bool
    utm: Dict[str, str] = Field(default_factory=dict)
    visitor_id: Optional[str] = Field(default=None, max_length=64)


def _eventbrite_code(ticket: dict, code: Optional[str]) -> Optional[str]:
    """An unrecognised code on an Eventbrite pass is passed to Eventbrite's checkout, which validates it."""
    if ticket.get("provider") == "eventbrite" and code and re.fullmatch(r"[A-Za-z0-9_-]{2,40}", code.strip()):
        return code.strip().upper()
    return None


def _registration_view(r: dict) -> dict:
    keep = ("code", "status", "ticket_id", "ticket_name", "first_name", "last_name", "email", "company", "job_title", "days", "dinners",
            "interests", "hackathon", "amount", "currency", "provider", "payment", "created_at", "event")
    return {k: r.get(k) for k in keep}


def _payment_instructions(event: dict, ticket: dict, reg: dict) -> Optional[dict]:
    provider = ticket.get("provider")
    if reg["amount"] <= 0 or provider == "free":
        return None
    if provider == "eventbrite":
        return {"provider": "eventbrite", "eventbrite_event_id": ticket.get("eventbrite_event_id") or event.get("eventbrite_event_id"), "promo_code": reg.get("promo_code")}
    if provider == "stripe_link" and ticket.get("payment_link"):
        link = ticket["payment_link"]
        sep = "&" if "?" in link else "?"
        from urllib.parse import quote as q
        return {"provider": "stripe_link", "url": f"{link}{sep}client_reference_id={reg['code']}&prefilled_email={q(reg['email'])}"}
    return {"provider": "invoice"}


@public.post("/{slug}/registrations", status_code=201)
async def register(slug: str, body: RegistrationIn, request: Request):
    if _rate_limited(f"reg:{_ip(request)}", 10):
        raise HTTPException(status_code=429, detail="Too many registrations from this network. Try again in a minute.")
    if not body.accept_terms:
        raise HTTPException(status_code=422, detail="Please accept the event terms and privacy notice.")
    event = await get_event(slug)
    if not event.get("registration_open"):
        raise HTTPException(status_code=409, detail="Registration is closed for this event.")
    ticket = next((t for t in event.get("tickets") or [] if t["id"] == body.ticket_id and t.get("active")), None)
    if not ticket:
        raise HTTPException(status_code=422, detail="Please choose an available pass.")
    if _sales_ended(ticket):
        raise HTTPException(status_code=409, detail="Sales for this pass have ended.")
    email = body.email.lower()
    existing = await db.event_registrations.find_one({"event": slug, "email": email, "status": {"$ne": "cancelled"}}, {"_id": 0})
    if existing:
        return {**_registration_view(existing), "already_registered": True, "payment": existing.get("payment")}

    counts = await _counts(slug)
    full = (event.get("capacity") and counts["_total"] >= event["capacity"]) or (ticket.get("capacity") and counts.get(ticket["id"], 0) >= ticket["capacity"])
    promo = _find_promo(event, body.promo_code, ticket["id"])
    price = _price(ticket, promo)
    day_ids = {d["id"] for d in event.get("days") or []}
    reg = {
        "id": str(uuid.uuid4()), "code": _new_code(), "event": slug, "event_name": event["name"],
        "ticket_id": ticket["id"], "ticket_name": ticket["name"], "provider": ticket.get("provider", "free"),
        "first_name": body.first_name.strip(), "last_name": body.last_name.strip(), "email": email,
        "company": body.company.strip(), "job_title": body.job_title.strip(), "phone": body.phone, "country": body.country,
        "interests": [i for i in dict.fromkeys(body.interests) if i in INTERESTS],
        "days": [d for d in dict.fromkeys(body.days) if d in day_ids],
        "dinners": [d for d in dict.fromkeys(body.dinners) if d in day_ids],
        "hackathon": body.hackathon, "dietary": body.dietary, "accessibility": body.accessibility, "how_heard": body.how_heard,
        "promo_code": promo["code"] if promo else _eventbrite_code(ticket, body.promo_code), "amount": price["total"], "list_price": price["price"], "discount": price["discount"],
        "currency": price["currency"], "billing_contact": body.billing_contact, "po_number": body.po_number,
        "marketing_opt_in": body.marketing_opt_in, "utm": {k: str(v)[:120] for k, v in list(body.utm.items())[:6]},
        "checked_in": False, "notes": "", "created_at": now_iso(), "updated_at": now_iso(),
    }
    if full:
        reg["status"] = "waitlisted"
    elif price["total"] <= 0 or reg["provider"] == "free":
        reg["status"] = "confirmed"
    elif reg["provider"] == "invoice" or (reg["provider"] == "stripe_link" and not ticket.get("payment_link")):
        reg["status"] = "invoice_requested"
    else:
        reg["status"] = "pending_payment"
    reg["payment"] = _payment_instructions(event, ticket, reg) if reg["status"] == "pending_payment" else None
    await db.event_registrations.insert_one(dict(reg))
    if promo:
        await db.event_configs.update_one({"slug": slug, "promo_codes.code": promo["code"]}, {"$inc": {"promo_codes.$.uses": 1}})
    cache.bump("events")

    # The same person is a sales signal: form inbox, lead alert and intent scoring.
    labels = [INTERESTS[i]["label"] for i in reg["interests"]]
    sub = {
        "id": str(uuid.uuid4()), "type": "event", "email": email, "name": f"{reg['first_name']} {reg['last_name']}",
        "company": reg["company"], "job_title": reg["job_title"], "phone": reg["phone"], "interest": ", ".join(labels) or None,
        "resource": event["name"], "message": f"{reg['ticket_name']} · {reg['status'].replace('_', ' ')} · code {reg['code']}",
        "source": "empower", "source_page": f"empower:{slug}", "created_at": now_iso(),
    }
    await db.submissions.insert_one(dict(sub))
    asyncio.create_task(notify_lead(sub))
    try:
        topics = [t for i in reg["interests"] for t in INTERESTS[i]["topics"]]
        await record_submission(sub, visitor_id=body.visitor_id, topics=topics, extra={"country": body.country}, tags=[slug])
    except Exception:
        logger.exception("lead scoring failed for event registration %s", reg["code"])
    asyncio.create_task(_send_confirmation(event, reg))
    return {**_registration_view(reg), "already_registered": False}


class LookupIn(BaseModel):
    email: EmailStr


@public.post("/{slug}/registrations/{code}/lookup")
async def lookup(slug: str, code: str, body: LookupIn, request: Request):
    if _rate_limited(f"lookup:{_ip(request)}", 20):
        raise HTTPException(status_code=429, detail="Too many attempts. Try again in a minute.")
    reg = await db.event_registrations.find_one({"event": slug, "code": code.upper(), "email": body.email.lower()}, {"_id": 0})
    if not reg:
        raise HTTPException(status_code=404, detail="We couldn't find that registration. Check the code and email.")
    return _registration_view(reg)


class PaymentReport(BaseModel):
    email: EmailStr
    provider: Literal["eventbrite", "stripe_link"]
    reference: Optional[str] = Field(default=None, max_length=120)


@public.post("/{slug}/registrations/{code}/payment")
async def report_payment(slug: str, code: str, body: PaymentReport):
    reg = await db.event_registrations.find_one({"event": slug, "code": code.upper(), "email": body.email.lower()}, {"_id": 0})
    if not reg:
        raise HTTPException(status_code=404, detail="Registration not found")
    if reg["status"] == "pending_payment":
        await db.event_registrations.update_one({"id": reg["id"]}, {"$set": {"status": "payment_reported", "payment_reference": body.reference, "payment_reported_at": now_iso(), "updated_at": now_iso()}})
        reg["status"] = "payment_reported"
        cache.bump("events")
    return _registration_view(reg)


async def _send_confirmation(event: dict, reg: dict) -> None:
    if not email_provider():
        return
    status_line = {
        "confirmed": "Your place is confirmed.",
        "waitlisted": "The event is full, so you're on the waitlist. We'll email you if a seat opens.",
        "invoice_requested": "We've reserved your place. Our events team will send an invoice to complete your registration.",
        "pending_payment": "Complete your payment to confirm your place.",
    }.get(reg["status"], "")
    html = (
        '<div style="font-family:Arial,Helvetica,sans-serif;max-width:560px;margin:0 auto;color:#0f172a">'
        f'<p style="color:#ED2423;font-size:11px;letter-spacing:2px;text-transform:uppercase">{escape(event["name"])}</p>'
        f'<h1 style="font-size:22px">Thanks for registering, {escape(reg["first_name"])}.</h1>'
        f'<p>{escape(status_line)}</p>'
        f'<p><strong>Registration code:</strong> {escape(reg["code"])}<br><strong>Pass:</strong> {escape(reg["ticket_name"])}<br>'
        f'<strong>When:</strong> October 28-30, 2026<br><strong>Where:</strong> {escape(event.get("venue") or "")}, {escape(event.get("address") or "")}</p>'
        f'<p style="color:#64748b;font-size:12px">Questions? Reply to this email or write to {escape(event.get("contact_email") or "info@solixempower.com")}. '
        'We never ask for passwords or card details by email.</p></div>'
    )
    try:
        await send_email(to=reg["email"], subject=f"Your {event['name']} registration ({reg['code']})", html=html, reply_to=event.get("contact_email"))
        await db.event_registrations.update_one({"id": reg["id"]}, {"$set": {"confirmation_sent_at": now_iso()}})
    except Exception:
        logger.exception("confirmation email failed for %s", reg["code"])


# --- Admin ------------------------------------------------------------------

@admin.get("")
async def list_events():
    for slug in DEFAULT_EVENTS:
        await get_event(slug)
    out = []
    async for e in db.event_configs.find({}, {"_id": 0, "promo_codes": 0}).sort("starts_at", -1):
        counts = await _counts(e["slug"])
        out.append({**e, "registered": counts["_total"]})
    return out


class TicketIn(BaseModel):
    id: str = Field(pattern=r"^[a-z0-9-]{2,40}$")
    name: str = Field(min_length=2, max_length=80)
    description: str = Field(default="", max_length=400)
    price: int = Field(ge=0, le=10_000_00)  # cents
    currency: str = Field(default="USD", pattern=r"^[A-Z]{3}$")
    provider: Provider = "free"
    capacity: Optional[int] = Field(default=None, ge=1, le=100000)
    eventbrite_event_id: Optional[str] = Field(default=None, pattern=r"^\d{6,20}$")
    payment_link: Optional[str] = Field(default=None, max_length=300)
    sales_end_at: Optional[str] = Field(default=None, max_length=40)
    active: bool = True


class PromoIn(BaseModel):
    code: str = Field(pattern=r"^[A-Za-z0-9_-]{3,40}$")
    percent_off: int = Field(ge=1, le=100)
    active: bool = True
    max_uses: Optional[int] = Field(default=None, ge=1)
    tickets: List[str] = Field(default_factory=list)
    uses: int = 0


class EventSettingsIn(BaseModel):
    name: str = Field(min_length=3, max_length=120)
    theme: str = Field(default="", max_length=300)
    registration_open: bool
    capacity: Optional[int] = Field(default=None, ge=1, le=100000)
    contact_email: Optional[EmailStr] = None
    eventbrite_event_id: Optional[str] = Field(default=None, pattern=r"^\d{6,20}$")
    refund_policy: Optional[str] = Field(default=None, max_length=400)
    tickets: List[TicketIn] = Field(min_length=1, max_length=20)
    promo_codes: List[PromoIn] = Field(default_factory=list, max_length=100)


@admin.get("/{slug}")
async def admin_event(slug: str):
    return await get_event(slug)


@admin.put("/{slug}", dependencies=[can_configure])
async def update_event(slug: str, body: EventSettingsIn, user: dict = Depends(get_current_admin)):
    await get_event(slug)
    data = body.model_dump()
    ids = [t["id"] for t in data["tickets"]]
    if len(ids) != len(set(ids)):
        raise HTTPException(status_code=422, detail="Each pass needs a unique ID.")
    for t in data["tickets"]:
        if t["provider"] == "stripe_link" and not (t.get("payment_link") or "").startswith("https://"):
            raise HTTPException(status_code=422, detail=f"'{t['name']}' uses a Stripe Payment Link: paste the https:// link.")
        if t["provider"] == "eventbrite" and not (t.get("eventbrite_event_id") or data.get("eventbrite_event_id")):
            raise HTTPException(status_code=422, detail=f"'{t['name']}' uses Eventbrite: add the Eventbrite event ID.")
        if t.get("sales_end_at"):
            try:
                datetime.fromisoformat(t["sales_end_at"])
            except ValueError:
                raise HTTPException(status_code=422, detail=f"'{t['name']}' sales end date must look like 2026-10-28T23:59:00-07:00.")
        if t["provider"] != "free" and t["price"] <= 0:
            raise HTTPException(status_code=422, detail=f"'{t['name']}' has a payment method but no price. Set a price or make it free.")
    codes = [p["code"].upper() for p in data["promo_codes"]]
    if len(codes) != len(set(codes)):
        raise HTTPException(status_code=422, detail="Promo codes must be unique.")
    for p in data["promo_codes"]:
        p["code"] = p["code"].upper()
    await db.event_configs.update_one({"slug": slug}, {"$set": {**data, "updated_at": now_iso(), "updated_by": user["email"]}})
    cache.bump("events")
    return await get_event(slug)


def _reg_query(slug: str, q: Optional[str], status: Optional[str], ticket: Optional[str], interest: Optional[str], checked_in: Optional[bool]) -> dict:
    query: dict = {"event": slug}
    if status and status != "all":
        query["status"] = {"$in": list(ACTIVE)} if status == "active" else status
    if ticket and ticket != "all":
        query["ticket_id"] = ticket
    if interest and interest != "all":
        query["interests"] = interest
    if checked_in is not None:
        query["checked_in"] = checked_in
    if q:
        rx = {"$regex": re.escape(q.strip()), "$options": "i"}
        query["$or"] = [{"first_name": rx}, {"last_name": rx}, {"email": rx}, {"company": rx}, {"code": rx}, {"job_title": rx}]
    return query


@admin.get("/{slug}/registrations")
async def registrations(slug: str, q: Optional[str] = None, status: Optional[str] = None, ticket: Optional[str] = None, interest: Optional[str] = None,
                        checked_in: Optional[bool] = None, page: int = Query(1, ge=1), page_size: int = Query(50, ge=1, le=200)):
    event = await get_event(slug)
    query = _reg_query(slug, q, status, ticket, interest, checked_in)
    total = await db.event_registrations.count_documents(query)
    items = await db.event_registrations.find(query, {"_id": 0}).sort("created_at", -1).skip((page - 1) * page_size).limit(page_size).to_list(page_size)
    everyone = await db.event_registrations.find({"event": slug}, {"_id": 0, "status": 1, "ticket_id": 1, "interests": 1, "days": 1, "dinners": 1,
                                                                  "checked_in": 1, "amount": 1, "currency": 1, "created_at": 1, "utm": 1, "hackathon": 1}).to_list(50000)
    active = [r for r in everyone if r["status"] in ACTIVE]
    by = lambda key: dict(sorted({k: sum(1 for r in active if k in (r.get(key) or [])) for k in {x for r in active for x in (r.get(key) or [])}}.items()))
    trend: Dict[str, int] = {}
    for r in everyone:
        trend[r["created_at"][:10]] = trend.get(r["created_at"][:10], 0) + 1
    sources: Dict[str, int] = {}
    for r in active:
        s = (r.get("utm") or {}).get("utm_source") or "direct"
        sources[s] = sources.get(s, 0) + 1
    stats = {
        "registered": len(active), "checked_in": sum(1 for r in active if r.get("checked_in")),
        "by_status": {s: sum(1 for r in everyone if r["status"] == s) for s in STATUSES},
        "by_ticket": {t["id"]: sum(1 for r in active if r["ticket_id"] == t["id"]) for t in event.get("tickets") or []},
        "by_interest": by("interests"), "by_day": by("days"), "dinners": by("dinners"),
        "hackathon": sum(1 for r in active if r.get("hackathon")),
        "revenue": sum(r.get("amount") or 0 for r in everyone if r["status"] in ("paid", "confirmed", "payment_reported")),
        "seats_left": max(0, event["capacity"] - len(active)) if event.get("capacity") else None,
        "trend": [{"date": d, "registrations": n} for d, n in sorted(trend.items())[-60:]],
        "sources": sorted(({"source": k, "count": v} for k, v in sources.items()), key=lambda x: -x["count"])[:8],
    }
    return {"items": items, "total": total, "page": page, "page_size": page_size, "stats": stats, "interests": {k: v["label"] for k, v in INTERESTS.items()}}


class RegistrationPatch(BaseModel):
    status: Optional[Literal[STATUSES]] = None  # type: ignore[valid-type]
    checked_in: Optional[bool] = None
    notes: Optional[str] = Field(default=None, max_length=2000)


@admin.patch("/{slug}/registrations/{reg_id}", dependencies=[can_work])
async def patch_registration(slug: str, reg_id: str, body: RegistrationPatch, user: dict = Depends(get_current_admin)):
    changes = {k: v for k, v in body.model_dump().items() if v is not None}
    if body.checked_in is not None:
        changes["checked_in_at"] = now_iso() if body.checked_in else None
    changes["updated_at"] = now_iso()
    changes["updated_by"] = user["email"]
    res = await db.event_registrations.find_one_and_update({"event": slug, "id": reg_id}, {"$set": changes}, projection={"_id": 0}, return_document=True)
    if not res:
        raise HTTPException(status_code=404, detail="Registration not found")
    cache.bump("events")
    return res


EXPORT = [("Code", "code"), ("Status", "status"), ("Pass", "ticket_name"), ("First name", "first_name"), ("Last name", "last_name"), ("Email", "email"),
          ("Company", "company"), ("Job title", "job_title"), ("Phone", "phone"), ("Country", "country"), ("Interests", "interests"), ("Days", "days"),
          ("Dinners", "dinners"), ("Hackathon", "hackathon"), ("Dietary", "dietary"), ("Accessibility", "accessibility"), ("Amount (cents)", "amount"),
          ("Currency", "currency"), ("Promo code", "promo_code"), ("Payment reference", "payment_reference"), ("PO number", "po_number"),
          ("Billing contact", "billing_contact"), ("How heard", "how_heard"), ("UTM source", "utm.utm_source"), ("UTM campaign", "utm.utm_campaign"),
          ("Checked in", "checked_in"), ("Registered", "created_at"), ("Notes", "notes")]


def _cell(r: dict, key: str):
    if key.startswith("utm."):
        v = (r.get("utm") or {}).get(key[4:])
    else:
        v = r.get(key)
    if isinstance(v, list):
        v = ", ".join(v)
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        v = "'" + v
    return "" if v is None else v


@admin.get("/{slug}/registrations-export")
async def export(slug: str, format: Literal["csv", "xlsx"] = "csv", q: Optional[str] = None, status: Optional[str] = None, ticket: Optional[str] = None,
                 interest: Optional[str] = None, checked_in: Optional[bool] = None):
    rows = await db.event_registrations.find(_reg_query(slug, q, status, ticket, interest, checked_in), {"_id": 0}).sort("created_at", 1).to_list(50000)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
    headers = [h for h, _ in EXPORT]
    data = [[_cell(r, k) for _, k in EXPORT] for r in rows]
    if format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Registrations"
        ws.append(headers)
        for row in data:
            ws.append(row)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0D192D")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{slug}-registrations-{stamp}.xlsx"'})
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(headers)
    w.writerows(data)
    return StreamingResponse(iter([out.getvalue()]), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="{slug}-registrations-{stamp}.csv"'})
