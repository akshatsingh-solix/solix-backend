"""Admin APIs for leads, leadership reports, exports, saved views, scoring
settings and staff users."""
from __future__ import annotations

import asyncio
import csv
import io
import re
import secrets
import uuid
from datetime import datetime, timedelta, timezone
from typing import Dict, List, Literal, Optional

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, EmailStr, Field

import cache
import scoring
from auth import ROLES, get_current_admin, hash_password, require_roles
from database import db, now_iso
from intent import get_scoring_settings, rescore_all, rollup_lead

router = APIRouter(prefix="/api/admin", tags=["leads-admin"], dependencies=[Depends(get_current_admin)])
can_work = Depends(require_roles("admin", "sales"))
admin_only = Depends(require_roles("admin"))

LINE_LABELS = {k: v["label"] for k, v in scoring.PRODUCT_LINES.items()}
CHANNELS = ("paid", "organic_search", "social", "email", "referral", "chat", "direct")
LEAD_LIST_FIELDS = {"_id": 0, "direct_scores": 0, "stage_history": 0, "visitor_ids": 0, "industries": 0}


# --- Filters ------------------------------------------------------------------

class LeadFilters(BaseModel):
    q: Optional[str] = None
    line: Optional[str] = None
    product: Optional[str] = None
    stage: Optional[str] = None
    owner: Optional[str] = None
    channel: Optional[str] = None
    country: Optional[str] = None
    industry: Optional[str] = None
    min_score: Optional[float] = None
    date_field: Literal["created_at", "last_activity_at", "mql_at"] = "created_at"
    date_from: Optional[str] = None
    date_to: Optional[str] = None
    active_days: Optional[int] = None


def filters_dep(
    q: Optional[str] = None, line: Optional[str] = None, product: Optional[str] = None, stage: Optional[str] = None,
    owner: Optional[str] = None, channel: Optional[str] = None, country: Optional[str] = None, industry: Optional[str] = None,
    min_score: Optional[float] = None, date_field: Literal["created_at", "last_activity_at", "mql_at"] = "created_at",
    date_from: Optional[str] = None, date_to: Optional[str] = None, active_days: Optional[int] = Query(None, ge=1, le=365),
) -> LeadFilters:
    return LeadFilters(**{k: v for k, v in locals().items()})


def _day_start(s: str) -> str:
    return s if "T" in s else f"{s}T00:00:00+00:00"


def _day_end(s: str) -> str:
    return s if "T" in s else f"{s}T23:59:59.999999+00:00"


def lead_query(f: LeadFilters) -> dict:
    clauses: list = []
    if f.q:
        rx = {"$regex": re.escape(f.q.strip()), "$options": "i"}
        clauses.append({"$or": [{"name": rx}, {"email": rx}, {"company": rx}, {"job_title": rx}, {"notes": rx}, {"tags": rx}]})
    if f.line and f.line != "all":
        clauses.append({"primary_line": None if f.line == "none" else f.line})
    if f.product and f.product != "all":
        clauses.append({f"product_scores.{f.product}": {"$gt": 0}})
    if f.stage and f.stage != "all":
        if f.stage == "mql_plus":
            clauses.append({"stage": {"$in": ["mql", "sal", "sql", "opportunity", "won"]}})
        elif f.stage == "sql_plus":
            clauses.append({"stage": {"$in": ["sql", "opportunity", "won"]}})
        else:
            clauses.append({"stage": f.stage})
    if f.owner and f.owner != "all":
        clauses.append({"owner": {"$in": [None, ""]}} if f.owner == "unassigned" else {"owner": f.owner})
    if f.channel and f.channel != "all":
        clauses.append({"channel": f.channel})
    if f.country and f.country != "all":
        clauses.append({"country": f.country})
    if f.industry and f.industry != "all":
        clauses.append({"industry": f.industry})
    if f.min_score is not None:
        clauses.append({"score": {"$gte": f.min_score}})
    if f.date_from:
        clauses.append({f.date_field: {"$gte": _day_start(f.date_from)}})
    if f.date_to:
        clauses.append({f.date_field: {"$lte": _day_end(f.date_to)}})
    if f.active_days:
        since = (datetime.now(timezone.utc) - timedelta(days=f.active_days)).isoformat()
        clauses.append({"last_activity_at": {"$gte": since}})
    return {"$and": clauses} if clauses else {}


SORTS = {"score": [("score", -1)], "recent": [("last_activity_at", -1)], "created": [("created_at", -1)], "mql": [("mql_at", -1)], "name": [("name", 1)]}


def _decorate(lead: dict) -> dict:
    lead["primary_line_label"] = LINE_LABELS.get(lead.get("primary_line"))
    lead.setdefault("stage", "lead")
    return lead


# --- Leads --------------------------------------------------------------------

@router.get("/leads")
async def list_leads(f: LeadFilters = Depends(filters_dep), sort: str = "score", page: int = Query(1, ge=1), page_size: int = Query(25, ge=1, le=200)):
    query = lead_query(f)
    total = await db.leads.count_documents(query)
    cursor = db.leads.find(query, LEAD_LIST_FIELDS)
    for field, direction in SORTS.get(sort, SORTS["score"]):
        cursor = cursor.sort(field, direction)
    items = await cursor.skip((page - 1) * page_size).limit(page_size).to_list(page_size)
    return {"items": [_decorate(i) for i in items], "total": total, "page": page, "page_size": page_size}


@router.get("/leads/meta")
async def leads_meta():
    """Everything the filter bar needs: lines, products, stages, owners, countries."""
    team = await db.settings.find_one({"key": "team"}, {"_id": 0}) or {}
    countries = sorted({c for c in await db.leads.distinct("country") if c})
    return {
        "lines": [{"key": k, "label": v["label"], "products": v["products"]} for k, v in scoring.PRODUCT_LINES.items()],
        "stages": scoring.STAGES, "channels": list(CHANNELS), "industries": sorted(scoring.INDUSTRIES),
        "owners": team.get("members", []), "countries": countries, "settings": await get_scoring_settings(),
    }


@router.get("/leads/{lead_id}")
async def get_lead(lead_id: str):
    lead = await db.leads.find_one({"id": lead_id}, {"_id": 0})
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found")
    subs = await db.submissions.find({"$or": [{"lead_id": lead["id"]}, {"email": lead["email"]}]}, {"_id": 0}).sort("created_at", -1).to_list(100)
    events = []
    if lead.get("visitor_ids"):
        events = await db.events.find({"visitor_id": {"$in": lead["visitor_ids"]}}, {"_id": 0, "at_dt": 0}).sort("at", -1).to_list(300)
    timeline = [{"kind": "form", "at": s["created_at"], "type": s["type"], "detail": s.get("interest") or s.get("resource") or s.get("message"), "source_page": s.get("source_page")} for s in subs]
    timeline += [{"kind": "event", "at": e["at"], "type": e["type"], "path": e.get("path"), "topics": e.get("topics"), "points": e.get("points"), "meta": e.get("meta")} for e in events]
    timeline += [{"kind": "stage", "at": h["at"], "type": h["stage"], "detail": h.get("reason"), "by": h.get("by")} for h in lead.get("stage_history") or []]
    deliveries = await db.deliveries.find({"email": lead["email"], "status": {"$in": ["sent", "failed"]}}, {"_id": 0}).sort("created_at", -1).to_list(50)
    timeline += [{"kind": "email", "at": d["created_at"], "type": f"asset_{d['status']}", "detail": d.get("title"), "meta": {"attached": d.get("attached"), "reason": d.get("detail")}} for d in deliveries]
    timeline.sort(key=lambda x: x["at"], reverse=True)
    products = sorted(((p, v) for p, v in (lead.get("product_scores") or {}).items()), key=lambda x: -x[1])
    return {
        "lead": _decorate(lead),
        "breakdown": {
            "fit": lead.get("fit_breakdown") or {},
            "lines": [{"key": k, "label": LINE_LABELS.get(k, k), "score": v} for k, v in sorted((lead.get("line_scores") or {}).items(), key=lambda x: -x[1])],
            "products": [{"slug": p, "line": scoring.PRODUCT_TO_LINE.get(p), "score": v} for p, v in products],
        },
        "submissions": subs,
        "timeline": timeline[:300],
    }


class LeadPatch(BaseModel):
    stage: Optional[Literal[tuple(scoring.STAGES)]] = None  # type: ignore[valid-type]
    owner: Optional[str] = Field(default=None, max_length=200)
    notes: Optional[str] = Field(default=None, max_length=5000)
    tags: Optional[List[str]] = Field(default=None, max_length=20)
    reason: Optional[str] = Field(default=None, max_length=300)


async def _apply_patch(lead: dict, body: LeadPatch, user: dict) -> dict:
    changes: dict = {}
    if body.owner is not None:
        changes["owner"] = body.owner.strip().lower() or None
    if body.notes is not None:
        changes["notes"] = body.notes
    if body.tags is not None:
        changes["tags"] = [t.strip()[:40] for t in body.tags if t.strip()]
    if body.stage and body.stage != lead.get("stage"):
        now = now_iso()
        changes["stage"] = body.stage
        changes["stage_reason"] = body.reason or f"Set by {user['email']}"
        changes["stage_changed_at"] = now
        if body.stage == "mql" and not lead.get("mql_at"):
            changes["mql_at"] = now
            changes["mql_line"] = lead.get("primary_line")
        if body.stage in ("sql", "opportunity", "won") and not lead.get("sql_at"):
            changes["sql_at"] = now
        changes["stage_history"] = ((lead.get("stage_history") or []) + [{"stage": body.stage, "at": now, "by": user["email"], "reason": body.reason}])[-50:]
    if changes:
        changes["updated_at"] = now_iso()
        await db.leads.update_one({"id": lead["id"]}, {"$set": changes})
        cache.bump("reports")
    return {**lead, **changes}


@router.patch("/leads/{lead_id}", dependencies=[can_work])
async def patch_lead(lead_id: str, body: LeadPatch, user: dict = Depends(get_current_admin)):
    lead = await db.leads.find_one({"id": lead_id}, {"_id": 0})
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found")
    return _decorate(await _apply_patch(lead, body, user))


class BulkPatch(LeadPatch):
    ids: List[str] = Field(min_length=1, max_length=500)


@router.post("/leads/bulk", dependencies=[can_work])
async def bulk_patch(body: BulkPatch, user: dict = Depends(get_current_admin)):
    n = 0
    async for lead in db.leads.find({"id": {"$in": body.ids}}, {"_id": 0}):
        await _apply_patch(lead, body, user)
        n += 1
    return {"updated": n}


@router.post("/leads/{lead_id}/rescore", dependencies=[can_work])
async def rescore_lead(lead_id: str):
    lead = await rollup_lead(lead_id)
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found")
    return _decorate({k: v for k, v in lead.items() if k not in ("direct_scores",)})


# --- Export -------------------------------------------------------------------

EXPORT_COLUMNS = [
    ("Name", "name"), ("Email", "email"), ("Company", "company"), ("Job title", "job_title"), ("Country", "country"),
    ("Company size", "company_size"), ("Industry", "industry"), ("Stage", "stage"), ("Stage reason", "stage_reason"),
    ("Primary product line", "primary_line_label"), ("Primary product", "primary_product"), ("Lead score", "score"),
    ("Fit score", "fit_score"), ("Behaviour score", "behaviour_score"),
] + [(f"Score: {v['label']}", f"line:{k}") for k, v in scoring.PRODUCT_LINES.items()] + [
    ("Channel", "channel"), ("UTM source", "ft:utm_source"), ("UTM medium", "ft:utm_medium"), ("UTM campaign", "ft:utm_campaign"),
    ("Landing page", "ft:landing"), ("Owner", "owner"), ("Forms submitted", "submissions_count"), ("Form types", "submission_types"),
    ("Pages viewed", "pages_viewed"), ("Sessions", "sessions"), ("Created", "created_at"), ("MQL at", "mql_at"),
    ("Last activity", "last_activity_at"), ("Tags", "tags"), ("Notes", "notes"),
]


def _cell(lead: dict, key: str):
    if key.startswith("line:"):
        return (lead.get("line_scores") or {}).get(key[5:], 0)
    if key.startswith("ft:"):
        return (lead.get("first_touch") or {}).get(key[3:]) or ""
    v = lead.get(key)
    if isinstance(v, list):
        return ", ".join(str(x) for x in v)
    return "" if v is None else v


@router.get("/leads-export")
async def export_leads(f: LeadFilters = Depends(filters_dep), format: Literal["csv", "xlsx"] = "csv", sort: str = "score"):
    cursor = db.leads.find(lead_query(f), LEAD_LIST_FIELDS)
    for field, direction in SORTS.get(sort, SORTS["score"]):
        cursor = cursor.sort(field, direction)
    leads = [_decorate(l) for l in await cursor.to_list(20000)]
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
    headers = [h for h, _ in EXPORT_COLUMNS]
    rows = [[_cell(l, k) for _, k in EXPORT_COLUMNS] for l in leads]
    if format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        wb = Workbook()
        ws = wb.active
        ws.title = "Leads"
        ws.append(headers)
        for r in rows:
            ws.append([_xl_safe(c) for c in r])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0D192D")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(40, max(12, len(h) + 2))
        summary = wb.create_sheet("Summary by product line")
        summary.append(["Product line", "Leads", "MQL or later", "SQL or later"])
        for key, label in LINE_LABELS.items():
            in_line = [l for l in leads if l.get("primary_line") == key]
            summary.append([label, len(in_line), sum(1 for l in in_line if l["stage"] not in ("lead", "disqualified")), sum(1 for l in in_line if l["stage"] in ("sql", "opportunity", "won"))])
        for c in summary[1]:
            c.font = Font(bold=True)
        summary.column_dimensions["A"].width = 44
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="solix-leads-{stamp}.xlsx"'})
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(headers)
    for r in rows:
        w.writerow([_csv_safe(c) for c in r])
    return StreamingResponse(iter([out.getvalue()]), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="solix-leads-{stamp}.csv"'})


def _csv_safe(v):
    # Neutralise spreadsheet formula injection from visitor-typed fields.
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


_xl_safe = _csv_safe


# --- Reports ------------------------------------------------------------------

def _window(date_from: Optional[str], date_to: Optional[str], days: int) -> tuple[str, str, list[str]]:
    end = datetime.fromisoformat(_day_end(date_to)) if date_to else datetime.now(timezone.utc)
    start = datetime.fromisoformat(_day_start(date_from)) if date_from else (end - timedelta(days=days - 1)).replace(hour=0, minute=0, second=0, microsecond=0)
    n = min(400, max(1, (end.date() - start.date()).days + 1))
    return start.isoformat(), end.isoformat(), [(start.date() + timedelta(days=i)).isoformat() for i in range(n)]


@router.get("/reports/overview")
async def overview(request: Request, date_from: Optional[str] = None, date_to: Optional[str] = None, days: int = Query(30, ge=1, le=365), line: Optional[str] = None):
    start, end, day_keys = _window(date_from, date_to, days)

    async def produce():
        lead_filter: dict = {"created_at": {"$gte": start, "$lte": end}}
        if line and line != "all":
            lead_filter["primary_line"] = line
        fields = {"_id": 0, "created_at": 1, "mql_at": 1, "sql_at": 1, "stage": 1, "primary_line": 1, "channel": 1, "owner": 1, "country": 1, "industry": 1, "score": 1, "primary_product": 1}
        mql_filter: dict = {"mql_at": {"$gte": start, "$lte": end}}
        if line and line != "all":
            mql_filter["mql_line"] = line
        # Independent queries run together; new visitors are counted per day in
        # the database instead of downloading every visitor record.
        leads, mqls, visitors_by_day, active_visitors = await asyncio.gather(
            db.leads.find(lead_filter, fields).to_list(50000),
            db.leads.find(mql_filter, {"_id": 0, "mql_at": 1, "mql_line": 1, "channel": 1, "owner": 1}).to_list(50000),
            db.visitors.aggregate([
                {"$match": {"created_at": {"$gte": start, "$lte": end}}},
                {"$group": {"_id": {"$substr": ["$created_at", 0, 10]}, "n": {"$sum": 1}}},
            ]).to_list(None),
            db.visitors.count_documents({"last_seen": {"$gte": start, "$lte": end}}),
        )

        new_visitors_total = sum(v["n"] for v in visitors_by_day)
        trend = {d: {"date": d, "visitors": 0, "leads": 0, "mqls": 0} for d in day_keys}
        for v in visitors_by_day:
            if v["_id"] in trend:
                trend[v["_id"]]["visitors"] += v["n"]
        for l in leads:
            if l["created_at"][:10] in trend:
                trend[l["created_at"][:10]]["leads"] += 1
        for m in mqls:
            if m["mql_at"][:10] in trend:
                trend[m["mql_at"][:10]]["mqls"] += 1

        by_line: Dict[str, dict] = {k: {"key": k, "label": v, "leads": 0, "mqls": 0, "sqls": 0} for k, v in LINE_LABELS.items()}
        by_line["none"] = {"key": "none", "label": "No product signal yet", "leads": 0, "mqls": 0, "sqls": 0}
        for l in leads:
            row = by_line[l.get("primary_line") or "none"]
            row["leads"] += 1
            if l.get("stage") in ("sql", "opportunity", "won"):
                row["sqls"] += 1
        for m in mqls:
            by_line[m.get("mql_line") or "none"]["mqls"] += 1

        by_channel = {c: {"key": c, "leads": 0, "mqls": 0} for c in CHANNELS}
        for l in leads:
            by_channel.setdefault(l.get("channel") or "direct", {"key": l.get("channel"), "leads": 0, "mqls": 0})["leads"] += 1
        for m in mqls:
            by_channel.setdefault(m.get("channel") or "direct", {"key": m.get("channel"), "leads": 0, "mqls": 0})["mqls"] += 1

        order = ["lead", "mql", "sal", "sql", "opportunity", "won"]
        reached = {s: 0 for s in order}
        for l in leads:
            st = l.get("stage") or "lead"
            # lost / disqualified leads only count as having been leads.
            for s in order[: order.index(st) + 1] if st in order else ["lead"]:
                reached[s] += 1
        funnel = [{"stage": "visitors", "count": new_visitors_total}] + [{"stage": s, "count": reached[s]} for s in order]

        owners: Dict[str, dict] = {}
        for l in leads:
            o = l.get("owner") or "unassigned"
            row = owners.setdefault(o, {"owner": o, "leads": 0, "mqls": 0, "sqls": 0})
            row["leads"] += 1
            if l.get("stage") not in ("lead", "disqualified", None):
                row["mqls"] += 1
            if l.get("stage") in ("sql", "opportunity", "won"):
                row["sqls"] += 1

        def top(key: str, n: int = 10):
            counts: Dict[str, int] = {}
            for l in leads:
                k = l.get(key)
                if k:
                    counts[k] = counts.get(k, 0) + 1
            return [{"key": k, "count": c} for k, c in sorted(counts.items(), key=lambda x: -x[1])[:n]]

        content_rows: Dict[str, dict] = {}
        async for e in db.events.find({"at": {"$gte": start, "$lte": end}, "type": {"$in": ["resource_view", "resource_download"]}, "path": {"$regex": "^/resources/"}}, {"_id": 0, "type": 1, "path": 1, "meta": 1}).limit(200000):
            row = content_rows.setdefault(e["path"], {"path": e["path"], "title": None, "views": 0, "downloads": 0})
            row["title"] = row["title"] or (e.get("meta") or {}).get("title")
            if e["type"] == "resource_download":
                row["downloads"] += 1
            else:
                row["views"] += 1
        page_rows: Dict[str, int] = {}
        async for e in db.events.find({"at": {"$gte": start, "$lte": end}, "type": "page_view"}, {"_id": 0, "path": 1}).limit(200000):
            page_rows[e["path"]] = page_rows.get(e["path"], 0) + 1

        # Product-level interest across all tracked visitors active in the window.
        interest: Dict[str, float] = {}
        async for v in db.visitors.find({"last_seen": {"$gte": start, "$lte": end}}, {"_id": 0, "product_scores": 1}).limit(100000):
            for p, s in (v.get("product_scores") or {}).items():
                interest[p] = round(interest.get(p, 0) + s, 1)

        total_leads = len(leads)
        mql_count = len(mqls)
        sql_count = sum(1 for l in leads if l.get("stage") in ("sql", "opportunity", "won"))
        return {
            "window": {"from": start, "to": end},
            "kpis": {
                "new_visitors": new_visitors_total, "active_visitors": active_visitors, "leads": total_leads, "mqls": mql_count,
                "sqls": sql_count, "won": sum(1 for l in leads if l.get("stage") == "won"),
                "visitor_to_lead": round(100 * total_leads / new_visitors_total, 1) if new_visitors_total else None,
                "lead_to_mql": round(100 * mql_count / total_leads, 1) if total_leads else None,
                "mql_to_sql": round(100 * sql_count / mql_count, 1) if mql_count else None,
                "unassigned_mqls": await db.leads.count_documents({"stage": "mql", "owner": {"$in": [None, ""]}}),
            },
            "trend": list(trend.values()),
            "by_line": [r for r in by_line.values() if r["leads"] or r["mqls"] or r["key"] != "none"],
            "by_channel": [r for r in by_channel.values()],
            "funnel": funnel,
            "owners": sorted(owners.values(), key=lambda r: -r["leads"]),
            "by_country": top("country"), "by_industry": top("industry"),
            "top_content": sorted(content_rows.values(), key=lambda r: -(r["views"] + 3 * r["downloads"]))[:10],
            "top_pages": [{"path": p, "views": c} for p, c in sorted(page_rows.items(), key=lambda x: -x[1])[:10]],
            "product_interest": [{"slug": p, "line": scoring.PRODUCT_TO_LINE.get(p), "score": s} for p, s in sorted(interest.items(), key=lambda x: -x[1])[:15]],
        }

    data = await cache.memo("reports", f"{start}|{end}|{line}", 60, produce)
    return cache.cached_json(request, data, max_age=30, swr=300, private=True)


# --- Saved views ----------------------------------------------------------------

class ViewIn(BaseModel):
    name: str = Field(min_length=1, max_length=80)
    filters: Dict[str, Optional[str]] = Field(default_factory=dict)
    sort: str = Field(default="score", max_length=20)


@router.get("/views")
async def list_views():
    return await db.lead_views.find({}, {"_id": 0}).sort("name", 1).to_list(200)


@router.post("/views", status_code=201)
async def create_view(body: ViewIn, user: dict = Depends(get_current_admin)):
    allowed = set(LeadFilters.model_fields)
    doc = {"id": str(uuid.uuid4()), "name": body.name.strip(), "filters": {k: v for k, v in body.filters.items() if k in allowed and v not in (None, "", "all")}, "sort": body.sort, "created_by": user["email"], "created_at": now_iso()}
    await db.lead_views.insert_one(dict(doc))
    doc.pop("_id", None)
    return doc


@router.delete("/views/{view_id}", status_code=204)
async def delete_view(view_id: str, user: dict = Depends(get_current_admin)):
    view = await db.lead_views.find_one({"id": view_id}, {"_id": 0})
    if not view:
        raise HTTPException(status_code=404, detail="View not found")
    if view["created_by"] != user["email"] and user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Only the creator or an admin can delete this view")
    await db.lead_views.delete_one({"id": view_id})


# --- Scoring settings -----------------------------------------------------------

class ScoringIn(BaseModel):
    mql_threshold: float = Field(ge=5, le=500)
    half_life_days: float = Field(ge=3, le=365)


@router.get("/scoring")
async def get_scoring():
    return {
        **await get_scoring_settings(), "event_weights": scoring.EVENT_WEIGHTS, "submission_weights": scoring.SUBMISSION_WEIGHTS,
        "hand_raise": sorted(scoring.HAND_RAISE), "lines": [{"key": k, "label": v["label"], "products": v["products"]} for k, v in scoring.PRODUCT_LINES.items()],
        "fit": {"business_email": 10, "seniority": "15 (director/VP/C-level) or 8 (manager/lead)", "company_size": "10 (1,000+), 6 (200+), 3 (50+)", "phone": 5},
    }


@router.put("/scoring", dependencies=[admin_only])
async def put_scoring(body: ScoringIn):
    await db.settings.update_one({"key": "scoring"}, {"$set": {**body.model_dump(), "updated_at": now_iso()}}, upsert=True)
    n = await rescore_all()
    cache.bump("reports")
    return {**await get_scoring(), "rescored": n}


# --- Staff users ----------------------------------------------------------------

class UserIn(BaseModel):
    email: EmailStr
    name: str = Field(min_length=1, max_length=120)
    role: Literal[ROLES]  # type: ignore[valid-type]


class UserPatch(BaseModel):
    name: Optional[str] = Field(default=None, min_length=1, max_length=120)
    role: Optional[Literal[ROLES]] = None  # type: ignore[valid-type]
    disabled: Optional[bool] = None
    reset_password: bool = False


def _temp_password() -> str:
    return secrets.token_urlsafe(9) + "!7a"


@router.get("/users", dependencies=[admin_only])
async def list_users():
    return await db.users.find({}, {"_id": 0, "password_hash": 0}).sort("created_at", 1).to_list(200)


@router.post("/users", status_code=201, dependencies=[admin_only])
async def create_user(body: UserIn):
    email = body.email.lower()
    if await db.users.find_one({"email": email}, {"_id": 1}):
        raise HTTPException(status_code=409, detail="A user with this email already exists")
    password = _temp_password()
    doc = {"id": str(uuid.uuid4()), "email": email, "name": body.name.strip(), "role": body.role, "password_hash": hash_password(password), "created_at": now_iso(), "disabled": False}
    await db.users.insert_one(dict(doc))
    cache.bump("users")
    return {"user": {k: v for k, v in doc.items() if k not in ("password_hash", "_id")}, "temporary_password": password}


@router.patch("/users/{user_id}", dependencies=[admin_only])
async def update_user(user_id: str, body: UserPatch, me: dict = Depends(get_current_admin)):
    target = await db.users.find_one({"id": user_id}, {"_id": 0})
    if not target:
        raise HTTPException(status_code=404, detail="User not found")
    if user_id == me["id"] and (body.disabled or (body.role and body.role != "admin")):
        raise HTTPException(status_code=400, detail="You can't remove your own admin access")
    changes = {k: v for k, v in body.model_dump(exclude={"reset_password"}).items() if v is not None}
    out: dict = {}
    if body.reset_password:
        password = _temp_password()
        changes["password_hash"] = hash_password(password)
        out["temporary_password"] = password
    if changes:
        await db.users.update_one({"id": user_id}, {"$set": changes})
        cache.bump("users")
    out["user"] = await db.users.find_one({"id": user_id}, {"_id": 0, "password_hash": 0})
    return out
