"""Admin API for the competitive-intent Radar (/api/admin/radar)."""
from __future__ import annotations

import asyncio
import io
import logging
from datetime import datetime, timedelta, timezone
from typing import Dict, List, Literal, Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import HTMLResponse, StreamingResponse
from pydantic import BaseModel, Field

import database
import radar_model as rm
import radar_service as rs
from auth import get_current_admin, require_roles
from database import now_iso

logger = logging.getLogger("solix.radar")
router = APIRouter(prefix="/api/admin/radar", tags=["radar"], dependencies=[Depends(get_current_admin)])
can_work = Depends(require_roles("admin", "sales"))
admin_only = Depends(require_roles("admin"))
_tasks: set = set()


def _db():
    return database.db


def _spawn(coro) -> None:
    t = asyncio.create_task(coro)
    _tasks.add(t)
    t.add_done_callback(_tasks.discard)


# --- Filters --------------------------------------------------------------------------

class RadarFilters(BaseModel):
    q: Optional[str] = None
    competitor: Optional[str] = None
    strength: Optional[str] = None
    country: Optional[str] = None
    state: Optional[str] = None
    territory: Optional[str] = None
    owner: Optional[str] = None
    status: Optional[str] = None
    line: Optional[str] = None
    tab: str = "all"
    ready_only: bool = False


def filters_dep(q: Optional[str] = None, competitor: Optional[str] = None, strength: Optional[str] = None, country: Optional[str] = None, state: Optional[str] = None,
                territory: Optional[str] = None, owner: Optional[str] = None, status: Optional[str] = None, line: Optional[str] = None, tab: str = "all", ready_only: bool = False) -> RadarFilters:
    return RadarFilters(**{k: v for k, v in locals().items()})


TAB_TYPES = {"jobs": {"competitor_job"}, "problems": {"problem_job"}, "linkedin": {"linkedin"}, "dns": {"dns", "cert"}, "news": {"news", "filing"}, "tenders": {"tender"}, "website": {"website"}}
LIST_FIELDS = {"_id": 0, "history": 0, "alias_keys": 0}


async def _visible(user: dict) -> dict:
    cfg = await _db().settings.find_one({"key": "radar"}, {"_id": 0, "reps_see_own_only": 1}) or {}
    if cfg.get("reps_see_own_only") and user.get("role") == "sales":
        return {"owner": user["email"].lower()}
    return {}


def _passes(c: dict, f: RadarFilters, fresh: str, *, with_tab: bool = True) -> bool:
    if f.q:
        q = f.q.strip().lower()
        hay = " ".join([c.get("name") or "", " ".join(c.get("aliases") or []), " ".join(c.get("domains") or []), " ".join(c.get("competitors") or []), c.get("owner_name") or ""]).lower()
        if q not in hay:
            return False
    if f.competitor and f.competitor not in (c.get("competitor_keys") or []):
        return False
    if f.strength and c.get("strength") != f.strength:
        return False
    if f.country and f.country not in (c.get("countries") or []) and (c.get("hq") or {}).get("country") != f.country:
        return False
    if f.state and f.state not in (c.get("states") or []):
        return False
    if f.territory and c.get("territory") != f.territory:
        return False
    if f.owner:
        if f.owner == "unassigned" and c.get("owner"):
            return False
        if f.owner != "unassigned" and c.get("owner") != f.owner:
            return False
    if f.status and c.get("status") != f.status:
        return False
    if f.line and c.get("primary_line") != f.line:
        return False
    if f.ready_only and (not c.get("research_id") or rm.verdict_class({"score": c.get("research_score"), "verdict": c.get("research_verdict")}) != "ready"):
        return False
    if not with_tab or f.tab in ("all", "territory"):
        return True
    if f.tab == "new":
        return fresh != "seen"
    if f.tab == "dives":
        return bool(c.get("research_id"))
    if f.tab == "ready":
        return c.get("status") == "ready"
    if f.tab == "surge":
        return bool(c.get("surge"))
    if f.tab == "excluded":
        return bool(c.get("excluded_reason"))
    want = TAB_TYPES.get(f.tab)
    return bool(want & set(c.get("signal_types") or [])) if want else True


def _sort_key(sort: str):
    if sort == "score":
        return lambda c: (-(c.get("score") or 0), c.get("name") or "")
    if sort == "recent":
        return lambda c: (str(c.get("latest_signal_at") or ""),)
    if sort == "name":
        return lambda c: ((c.get("name") or "").lower(),)
    if sort == "dives":
        return lambda c: (-(c.get("research_score") if c.get("research_score") is not None else -1), rm.STRENGTH_RANK.get(c.get("strength"), 3))
    return lambda c: (rm.STRENGTH_RANK.get(c.get("strength"), 3), -(c.get("research_score") if c.get("research_score") is not None else -1),
                      -len(c.get("signal_types") or []), -(c.get("signal_count") or 0), "".join(chr(255 - ord(ch)) for ch in str(c.get("latest_signal_at") or ""))[:30], c.get("name") or "")


async def _base(user: dict) -> list:
    return await _db().companies.find(await _visible(user), LIST_FIELDS).to_list(20000)


def _card(c: dict, fresh: str) -> dict:
    return {**c, "fresh": fresh, "verdict_class": rm.verdict_class({"score": c.get("research_score"), "verdict": c.get("research_verdict")}) if c.get("research_id") else None,
            "status_label": rm.STATUS_LABELS.get(c.get("status"), c.get("status"))}


# --- Reads ---------------------------------------------------------------------------------

@router.get("/meta")
async def meta():
    terr = await rs.territories()
    settings = await rs.radar_settings()
    return {
        "strengths": rm.STRENGTHS, "statuses": [{"key": s, "label": rm.STATUS_LABELS[s]} for s in rm.STATUSES], "territories": rm.TERRITORY_ORDER,
        "competitors": [{"key": c["key"], "product": c["product"], "vendor": c["vendor"]} for c in rm.COMPETITORS], "highlight": rm.highlight_terms(),
        "signal_types": rm.SIGNAL_TYPES, "team": await rs.team_members(), "reps": sorted({na["rep"] for na in terr.get("named_accounts", [])} | set((terr.get("us_states") or {}).values()) | set((terr.get("territory_owners") or {}).values())),
        "settings": {k: settings[k] for k in ("research_per_run", "refresh_days", "collectors")}, "ready_score": rm.READY_SCORE,
    }


@router.get("/summary")
async def summary(f: RadarFilters = Depends(filters_dep), user: dict = Depends(get_current_admin)):
    companies = await _base(user)
    seen = await rs.seen_map(user["id"])
    fresh = {c["id"]: rs.freshness(c, seen) for c in companies}
    base = [c for c in companies if _passes(c, f, fresh[c["id"]], with_tab=False)]
    ids = [c["id"] for c in base]
    today = datetime.now(timezone.utc).date().isoformat()
    comp_counts: Dict[str, int] = {}
    new_today = 0
    total_signals = 0
    async for s in _db().company_signals.find({"company_id": {"$in": ids}}, {"_id": 0, "competitor_product": 1, "competitor_key": 1, "first_seen_at": 1}):
        total_signals += 1
        if s.get("competitor_key"):
            label = (rm.COMPETITOR_BY_KEY.get(s["competitor_key"]) or {}).get("product") or s.get("competitor_product")
            comp_counts[label] = comp_counts.get(label, 0) + 1
        if str(s.get("first_seen_at") or "")[:10] == today:
            new_today += 1
    by_strength = {s: sum(1 for c in base if c.get("strength") == s) for s in rm.STRENGTHS}
    dives = [c for c in base if c.get("research_id")]
    ready = [c for c in dives if rm.verdict_class({"score": c.get("research_score"), "verdict": c.get("research_verdict")}) == "ready"]
    terr: Dict[str, dict] = {}
    for c in base:
        t = terr.setdefault(c.get("territory") or "Unrouted", {"territory": c.get("territory") or "Unrouted", "accounts": 0, "strong": 0, "reps": {}})
        t["accounts"] += 1
        t["strong"] += c.get("strength") == "Strong"
        rep = c.get("owner_name") or "Unassigned"
        r = t["reps"].setdefault(rep, {"rep": rep, "owner": c.get("owner"), "accounts": 0, "strong": 0})
        r["accounts"] += 1
        r["strong"] += c.get("strength") == "Strong"
    last = await _db().ops_runs.find_one({"job": "radar_daily"}, {"_id": 0, "notes": 0}, sort=[("started_at", -1)])
    visit = await _db().settings.find_one({"key": f"radar_visit:{user['id']}"}, {"_id": 0}) or {}
    await _db().settings.update_one({"key": f"radar_visit:{user['id']}"}, {"$set": {"last_visit": now_iso(), "prev_visit": visit.get("last_visit")}}, upsert=True)
    return {
        "accounts": len(base), "signals": total_signals, "by_strength": by_strength, "multi_signal": sum(1 for c in base if len(c.get("signal_types") or []) >= 2),
        "new_today": new_today, "deep_dives": len(dives), "ready": len(ready), "surging": sum(1 for c in base if c.get("surge")),
        "unassigned": sum(1 for c in base if not c.get("owner")), "excluded": sum(1 for c in base if c.get("excluded_reason")),
        "fresh": {"new": sum(1 for c in base if fresh[c["id"]] == "new"), "updated": sum(1 for c in base if fresh[c["id"]] == "updated")},
        "first_visit": not seen, "last_visit": visit.get("last_visit"),
        "competitors": [{"product": k, "signals": v} for k, v in sorted(comp_counts.items(), key=lambda kv: -kv[1])[:8]],
        "scope": [{"id": c["id"], "name": c["name"], "strength": c.get("strength"), "signals": c.get("signal_count") or 0, "dive": c.get("research_score"), "surge": c.get("surge"), "fresh": fresh[c["id"]]} for c in base],
        "territories": [{**t, "reps": sorted(t["reps"].values(), key=lambda r: -r["accounts"])} for t in sorted(terr.values(), key=lambda t: rm.TERRITORY_ORDER.index(t["territory"]) if t["territory"] in rm.TERRITORY_ORDER else 99)],
        "tabs": {tab: sum(1 for c in base if _passes(c, RadarFilters(tab=tab), fresh[c["id"]])) for tab in ("new", "dives", "jobs", "problems", "linkedin", "dns", "news", "tenders", "website", "ready", "surge", "excluded")},
        "last_run": last,
    }


@router.get("/accounts")
async def list_accounts(f: RadarFilters = Depends(filters_dep), sort: str = "strength", page: int = Query(1, ge=1), page_size: int = Query(40, ge=1, le=500), user: dict = Depends(get_current_admin)):
    companies = await _base(user)
    seen = await rs.seen_map(user["id"])
    rows = []
    for c in companies:
        fr = rs.freshness(c, seen)
        if _passes(c, f, fr):
            rows.append(_card(c, fr))
    if f.tab == "dives" and sort == "strength":
        sort = "dives"
    rows.sort(key=_sort_key(sort), reverse=(sort == "recent"))
    start = (page - 1) * page_size
    return {"items": rows[start:start + page_size], "total": len(rows), "page": page, "page_size": page_size}


@router.get("/accounts/{company_id}")
async def get_account(company_id: str, user: dict = Depends(get_current_admin)):
    c = await _db().companies.find_one({"id": company_id, **(await _visible(user))}, {"_id": 0})
    if not c:
        raise HTTPException(status_code=404, detail="Account not found")
    signals = await _db().company_signals.find({"company_id": company_id}, {"_id": 0, "raw": 0}).sort([("posted_at", -1), ("first_seen_at", -1)]).to_list(1000)
    signals.sort(key=lambda s: (rm.STRENGTH_RANK.get(s.get("role_strength") or "Unrated", 3) if s["type"] == "competitor_job" else 4, str(s.get("posted_at") or s.get("first_seen_at") or "")), reverse=False)
    research = await _db().radar_research.find({"company_id": company_id}, {"_id": 0}).sort("created_at", -1).to_list(10)
    leads = await _db().leads.find({"$or": [{"company_id": company_id}, {"id": {"$in": c.get("lead_ids") or []}}]}, {"_id": 0, "id": 1, "email": 1, "name": 1, "job_title": 1, "stage": 1, "score": 1, "primary_line": 1, "last_activity_at": 1, "owner": 1, "channel": 1}).to_list(100)
    contacts = await _db().contacts.find({"company_id": company_id}, {"_id": 0}).sort("created_at", 1).to_list(200)
    enrollments = await _db().enrollments.find({"company_id": company_id}, {"_id": 0}).sort("created_at", -1).to_list(100)
    for e in enrollments:  # the rendered copy stays out of the drawer payload; only its length is shown
        e["steps_total"] = len(e.pop("steps", None) or [])
    seen = await rs.seen_map(user["id"])
    return {"account": _card(c, rs.freshness(c, seen)), "signals": signals, "research": research, "leads": leads, "contacts": contacts, "enrollments": enrollments}


# --- Writes --------------------------------------------------------------------------------

class AccountPatch(BaseModel):
    owner: Optional[str] = Field(default=None, max_length=200)
    status: Optional[str] = None
    notes: Optional[str] = Field(default=None, max_length=5000)
    reason: Optional[str] = Field(default=None, max_length=300)
    is_customer: Optional[bool] = None
    revenue_band: Optional[str] = Field(default=None, max_length=40)
    employees: Optional[int] = Field(default=None, ge=0)
    industry: Optional[str] = Field(default=None, max_length=80)
    hq: Optional[Dict[str, Optional[str]]] = None
    domains: Optional[List[str]] = Field(default=None, max_length=20)
    manual_exclusion: Optional[str] = Field(default=None, max_length=200)


@router.patch("/accounts/{company_id}", dependencies=[can_work])
async def patch_account(company_id: str, body: AccountPatch, user: dict = Depends(get_current_admin)):
    c = await _db().companies.find_one({"id": company_id}, {"_id": 0})
    if not c:
        raise HTTPException(status_code=404, detail="Account not found")
    payload = body.model_dump(exclude_unset=True)
    try:
        updated = await rs.patch_company(c, payload, user)
    except ValueError as e:
        raise HTTPException(status_code=422, detail=str(e))
    if payload.get("status") in ("contacted", "meeting_booked", "won", "not_a_fit"):
        import outbound
        await outbound.on_account_status(updated, user)
    await rs.mark_seen(user["id"], [company_id])
    return updated


class AccountIn(BaseModel):
    name: str = Field(min_length=1, max_length=200)
    domain: Optional[str] = Field(default=None, max_length=200)
    country: Optional[str] = Field(default=None, max_length=80)
    state: Optional[str] = Field(default=None, max_length=80)
    zip: Optional[str] = Field(default=None, max_length=12)


@router.post("/accounts", status_code=201, dependencies=[can_work])
async def create_account(body: AccountIn, user: dict = Depends(get_current_admin)):
    c = await rs.get_or_create_company(body.name, domain=body.domain, hq={"country": body.country, "state": body.state, "zip": body.zip}, source="manual")
    await rs.add_history(c["id"], user["email"], "created", "Added by hand")
    return await rs.recompute(c["id"])


class SignalIn(BaseModel):
    type: Literal["linkedin", "news", "filing", "tender", "competitor_job", "problem_job", "dns"]
    title: str = Field(min_length=1, max_length=300)
    evidence_url: Optional[str] = Field(default=None, max_length=1000)
    evidence_text: Optional[str] = Field(default=None, max_length=2000)
    competitor_key: Optional[str] = None
    person: Optional[str] = Field(default=None, max_length=120)
    person_title: Optional[str] = Field(default=None, max_length=200)
    country: Optional[str] = Field(default=None, max_length=80)
    posted_at: Optional[str] = None


@router.post("/accounts/{company_id}/signals", status_code=201, dependencies=[can_work])
async def add_signal(company_id: str, body: SignalIn, user: dict = Depends(get_current_admin)):
    c = await _db().companies.find_one({"id": company_id}, {"_id": 0})
    if not c:
        raise HTTPException(status_code=404, detail="Account not found")
    if body.evidence_url and not body.evidence_url.startswith(("http://", "https://")):
        raise HTTPException(status_code=422, detail="Evidence must be a web address")
    sig = body.model_dump()
    title = sig.get("person_title") or sig.get("title")
    sig.update({"source": "manual", "role_title": title, "role_strength": rm.role_strength(title or ""), "posted_at": body.posted_at or now_iso()[:10],
                "competitor_product": (rm.COMPETITOR_BY_KEY.get(body.competitor_key) or {}).get("product"), "meta": {"added_by": user["email"]}})
    await rs.upsert_signal(c, sig)
    await rs.add_history(company_id, user["email"], "signal", f"Added {body.type.replace('_', ' ')}: {body.title}")
    return await rs.recompute(company_id)


@router.delete("/signals/{signal_id}", status_code=204, dependencies=[admin_only])
async def delete_signal(signal_id: str, user: dict = Depends(get_current_admin)):
    s = await _db().company_signals.find_one({"id": signal_id}, {"_id": 0})
    if not s:
        raise HTTPException(status_code=404, detail="Signal not found")
    await _db().company_signals.delete_one({"id": signal_id})
    await rs.add_history(s["company_id"], user["email"], "signal", f"Removed signal: {s.get('title') or s.get('type')}")
    await rs.recompute(s["company_id"])


class SeenIn(BaseModel):
    ids: Optional[List[str]] = Field(default=None, max_length=5000)
    all: bool = False


@router.post("/seen")
async def seen(body: SeenIn, user: dict = Depends(get_current_admin)):
    n = await rs.mark_seen(user["id"], None if body.all else (body.ids or []))
    return {"marked": n}


# --- Research ------------------------------------------------------------------------------

@router.post("/accounts/{company_id}/research", dependencies=[can_work])
async def request_research(company_id: str, user: dict = Depends(get_current_admin)):
    import radar_research
    c = await _db().companies.find_one({"id": company_id}, {"_id": 0})
    if not c:
        raise HTTPException(status_code=404, detail="Account not found")
    job = await radar_research.queue(c, by=user["email"])
    _spawn(radar_research.run_queue(limit=1, only=job["id"]))
    return job


@router.get("/research")
async def research_queue(status: Optional[str] = "draft", limit: int = Query(50, le=200)):
    q = {"status": status} if status and status != "all" else {}
    items = await _db().radar_research.find(q, {"_id": 0}).sort("created_at", -1).to_list(limit)
    names = {c["id"]: c["name"] async for c in _db().companies.find({"id": {"$in": [i["company_id"] for i in items]}}, {"_id": 0, "id": 1, "name": 1})}
    for i in items:
        i["company"] = names.get(i["company_id"], i.get("company"))
    return items


class ResearchPatch(BaseModel):
    action: Literal["approve", "reject", "save"]
    headline: Optional[str] = Field(default=None, max_length=500)
    summary: Optional[str] = Field(default=None, max_length=6000)
    trigger: Optional[str] = Field(default=None, max_length=3000)
    why_solix_wins: Optional[str] = Field(default=None, max_length=4000)
    gaps: Optional[str] = Field(default=None, max_length=3000)
    score: Optional[int] = Field(default=None, ge=0, le=10)
    verdict: Optional[str] = Field(default=None, max_length=40)
    first_touch: Optional[Dict[str, Optional[str]]] = None
    reason: Optional[str] = Field(default=None, max_length=300)


@router.patch("/research/{research_id}", dependencies=[can_work])
async def review_research(research_id: str, body: ResearchPatch, user: dict = Depends(get_current_admin)):
    import radar_research
    r = await _db().radar_research.find_one({"id": research_id}, {"_id": 0})
    if not r:
        raise HTTPException(status_code=404, detail="Research not found")
    edits = {k: v for k, v in body.model_dump(exclude={"action", "reason"}, exclude_none=True).items()}
    if edits:
        r.update(edits)
        r["checks"] = radar_research.validate(r, await _db().companies.find_one({"id": r["company_id"]}, {"_id": 0}) or {})
    sets = {**edits, "checks": r.get("checks", []), "updated_at": now_iso()}
    if body.action == "approve":
        blocking = [c for c in r.get("checks", []) if c.get("level") == "block"]
        if blocking:
            raise HTTPException(status_code=422, detail="Fix before approving: " + "; ".join(c["message"] for c in blocking))
        sets.update({"status": "approved", "reviewed_by": user["email"], "reviewed_at": now_iso()})
        await _db().radar_research.update_many({"company_id": r["company_id"], "status": "approved", "id": {"$ne": research_id}}, {"$set": {"status": "stale"}})
    elif body.action == "reject":
        sets.update({"status": "rejected", "reviewed_by": user["email"], "reviewed_at": now_iso(), "reject_reason": body.reason})
    await _db().radar_research.update_one({"id": research_id}, {"$set": sets})
    await rs.add_history(r["company_id"], user["email"], "research", {"approve": "Deep dive approved", "reject": "Deep dive rejected", "save": "Deep dive edited"}[body.action])
    await rs.recompute(r["company_id"])
    return await _db().radar_research.find_one({"id": research_id}, {"_id": 0})


# --- Promote to the lead funnel ----------------------------------------------------------------

class PromoteIn(BaseModel):
    email: str = Field(min_length=3, max_length=200)
    name: Optional[str] = Field(default=None, max_length=120)
    job_title: Optional[str] = Field(default=None, max_length=200)
    note: Optional[str] = Field(default=None, max_length=500)


@router.post("/accounts/{company_id}/promote", dependencies=[can_work])
async def promote(company_id: str, body: PromoteIn, user: dict = Depends(get_current_admin)):
    import outbound
    c = await _db().companies.find_one({"id": company_id}, {"_id": 0})
    if not c:
        raise HTTPException(status_code=404, detail="Account not found")
    lead = await outbound.promote_to_lead(c, {"email": body.email, "name": body.name, "job_title": body.job_title}, reason=body.note or f"Promoted from the Radar by {user['email']}", by=user["email"])
    return {"lead_id": lead["id"], "stage": lead.get("stage")}


# --- Configuration ---------------------------------------------------------------------------

class TerritoriesIn(BaseModel):
    territory_owners: Dict[str, str] = Field(default_factory=dict)
    us_states: Dict[str, str] = Field(default_factory=dict)
    ca_split: Dict[str, object] = Field(default_factory=dict)
    countries: Dict[str, str] = Field(default_factory=dict)
    named_accounts: List[Dict[str, str]] = Field(default_factory=list, max_length=2000)
    rep_territory: Dict[str, str] = Field(default_factory=dict)
    rep_emails: Dict[str, str] = Field(default_factory=dict)


@router.get("/territories")
async def get_territories():
    return await rs.territories()


@router.put("/territories", dependencies=[admin_only])
async def put_territories(body: TerritoriesIn, user: dict = Depends(get_current_admin)):
    values = body.model_dump()
    values["named_accounts"] = [{"match": (n.get("match") or "").strip().lower(), "rep": (n.get("rep") or "").strip()} for n in values["named_accounts"] if (n.get("match") or "").strip() and (n.get("rep") or "").strip()]
    values["rep_emails"] = {k.strip(): v.strip().lower() for k, v in values["rep_emails"].items() if k.strip() and v.strip()}
    saved = await rs.save_territories(values)
    _spawn(rs.recompute_all())
    return saved


@router.get("/routing/preview")
async def routing_preview(company: str, country: Optional[str] = None, location: Optional[str] = None, state: Optional[str] = None, zip: Optional[str] = None):
    terr = await rs.territories()
    r = rm.route(company, country, location, state, zip, terr)
    return {**r, "owner": rs.rep_email(r["rep"], terr, await rs.team_members())}


class ConfigIn(BaseModel):
    targets: Optional[List[Dict[str, Optional[str]]]] = Field(default=None, max_length=3000)
    dns_domains: Optional[List[str]] = Field(default=None, max_length=5000)
    collectors: Optional[Dict[str, bool]] = None
    credit_cap: Optional[int] = Field(default=None, ge=0, le=100000)
    research_per_run: Optional[int] = Field(default=None, ge=0, le=50)
    refresh_days: Optional[int] = Field(default=None, ge=1, le=365)
    min_employees: Optional[int] = Field(default=None, ge=0)
    countries: Optional[List[str]] = None
    tender_countries: Optional[List[str]] = None
    reps_see_own_only: Optional[bool] = None


@router.get("/config")
async def get_config():
    s = await rs.radar_settings()
    extra = await _db().settings.find_one({"key": "radar"}, {"_id": 0, "reps_see_own_only": 1}) or {}
    return {**s, "reps_see_own_only": bool(extra.get("reps_see_own_only"))}


@router.put("/config", dependencies=[admin_only])
async def put_config(body: ConfigIn):
    values = body.model_dump(exclude_none=True)
    if "dns_domains" in values:
        values["dns_domains"] = sorted({d for d in (rm.domain_of(x) for x in values["dns_domains"]) if d})
    if "targets" in values:
        values["targets"] = [t for t in values["targets"] if (t.get("company") or t.get("domain"))]
    if "reps_see_own_only" in values:
        await _db().settings.update_one({"key": "radar"}, {"$set": {"reps_see_own_only": values.pop("reps_see_own_only")}}, upsert=True)
    await rs.save_radar_settings(values)
    return await get_config()


# --- Import, export, runs ---------------------------------------------------------------------

@router.post("/import", dependencies=[admin_only])
async def import_radar(file: UploadFile = File(...), user: dict = Depends(get_current_admin)):
    raw = await file.read()
    if len(raw) > 25 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="File is larger than 25 MB")
    try:
        data = rs.parse_radar_payload(raw)
    except (ValueError, UnicodeDecodeError) as e:
        raise HTTPException(status_code=422, detail=f"Couldn't read that file: {e}")
    import ops
    async with ops.run("radar_import", by=user["email"]) as r:
        r["counts"] = await rs.import_radar(data, by=user["email"])
    return r["counts"]


@router.get("/legacy.json", dependencies=[can_work])
async def legacy(owner: Optional[str] = None):
    return await rs.legacy_export(owner)


@router.get("/export.xlsx", dependencies=[can_work])
async def export_xlsx(owner: Optional[str] = None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    data = await rs.legacy_export(owner)
    wb = Workbook()
    sheets = [("Summary", None), ("Job Postings", "jobs"), ("Problem Signals", "problems"), ("LinkedIn Engagers", "linkedin"), ("DNS Signals", "dns")]
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Company", "Lead Strength"])
    for a in data["accounts"]:
        ws.append([a["Company"], a["Lead Strength"]])
    for title, key in sheets[1:]:
        rows = data[key]
        sh = wb.create_sheet(title)
        headers = list(rows[0].keys()) if rows else ["Company"]
        sh.append(headers)
        for r in rows:
            sh.append([_xl(r.get(h)) for h in headers])
    runs = wb.create_sheet("Run Log")
    runs.append(["Job", "Started", "Finished", "Status", "Counts"])
    async for run in _db().ops_runs.find({"job": {"$regex": "^radar"}}, {"_id": 0}).sort("started_at", -1).limit(200):
        runs.append([run["job"], run.get("started_at"), run.get("finished_at"), run.get("status"), str(run.get("counts") or {})])
    for sh in wb.worksheets:
        for c in sh[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0D192D")
        sh.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="Competitive_Intent_Radar-{stamp}.xlsx"'})


def _xl(v):
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


class RunIn(BaseModel):
    steps: List[Literal["collect", "score", "research", "digest"]] = Field(default_factory=lambda: ["collect", "score", "research"])


@router.post("/run", dependencies=[admin_only])
async def run_now(body: RunIn, user: dict = Depends(get_current_admin)):
    import radar_run
    two_hours_ago = (datetime.now(timezone.utc) - timedelta(hours=2)).isoformat()
    if await _db().ops_runs.find_one({"job": "radar_daily", "status": "running", "started_at": {"$gte": two_hours_ago}}, {"_id": 1}):
        raise HTTPException(status_code=409, detail="A Radar run is already in progress")
    _spawn(radar_run.daily(tuple(body.steps), by=user["email"]))
    return {"started": True, "steps": body.steps}


@router.post("/recompute", dependencies=[admin_only])
async def recompute_all():
    return {"recomputed": await rs.recompute_all()}


@router.get("/digest/preview", response_class=HTMLResponse, dependencies=[can_work])
async def digest_preview(owner: str):
    import radar_run
    html, n = await radar_run.digest_html(owner)
    return HTMLResponse(html or "<p>Nothing new for this owner since the last digest.</p>")
